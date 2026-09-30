"""TC 라이브러리 저장소 — state/tc_library/{suite}/ (PRD §5, F5.2~F5.4, F5.11, F6.7).

파일 구성:
  cases.json             {"suite", "sheets": [...], "cases": [...]}  ← update_state 원자 쓰기
  history.jsonl          변경 1건 = 1줄 (append 전용)
  template.xlsx          가져온 원본의 사본 (내보내기 템플릿)
  template_profile.json  시트별 TemplateProfile
"""
from __future__ import annotations

import json
import secrets
import shutil
import os
import time
import threading
from contextlib import contextmanager
from functools import wraps
from pathlib import Path

import _paths
from _state import read_state, update_state
from _tc_model import EDITABLE_FIELDS, next_case_id, now_iso, validate_case
from _tc_template import TemplateProfile
from _validators import is_valid_group_name


class LibraryError(Exception):
    def __init__(self, message: str, code: str, status: int = 400):
        super().__init__(message)
        self.code = code
        self.status = status


class RevConflict(LibraryError):
    def __init__(self, server_case: dict):
        super().__init__("다른 곳에서 먼저 바뀌었습니다", "REV_CONFLICT", 409)
        self.server_case = server_case


# API 경로 조각과 겹치는 이름은 스위트로 쓸 수 없다 (/api/tc-library/{profiles|sources|jobs|…})
RESERVED_SUITES = {"import", "exports", "sources", "profiles", "jobs", "credentials", "source-diff", "trash"}
# 처음 접속할 때 쓰는 빈 양식 스위트. 삭제할 수 없다
DEFAULT_SUITE = "기본양식"


def suite_dir(suite: str) -> Path:
    if not suite or suite.startswith("_") or suite in RESERVED_SUITES or not is_valid_group_name(suite):
        raise LibraryError(f"스위트 이름이 올바르지 않습니다: {suite!r}", "INVALID_SUITE")
    return _paths.TC_LIBRARY_DIR / suite



# ponytail: serialize a suite; per-resource locks only if measured import throughput requires it.
_SUITE_LOCKS: dict[str, threading.RLock] = {}
_SUITE_LOCKS_GUARD = threading.Lock()
_SUITE_HELD = threading.local()


@contextmanager
def suite_lock(suite: str):
    """Reentrant thread/process lock; kernel releases the native lock after a process crash."""
    root = suite_dir(suite)
    key = str(root.resolve())
    held = getattr(_SUITE_HELD, 'keys', set())
    if key in held:
        yield
        return
    with _SUITE_LOCKS_GUARD:
        lock = _SUITE_LOCKS.setdefault(key, threading.RLock())
    with lock:
        lock_dir = _paths.TC_LIBRARY_DIR / '_locks'
        lock_dir.mkdir(parents=True, exist_ok=True)
        import hashlib
        lock_path = lock_dir / (hashlib.sha256(key.encode()).hexdigest() + '.lock')
        with lock_path.open('a+b') as file:
            if os.name == 'nt':
                import msvcrt
                if not file.seek(0, os.SEEK_END):
                    file.write(b'0')
                    file.flush()
                deadline = time.monotonic() + _paths.LOCK_TIMEOUT_SECS
                while True:
                    file.seek(0)
                    try:
                        msvcrt.locking(file.fileno(), msvcrt.LK_NBLCK, 1)
                        break
                    except OSError as exc:
                        if time.monotonic() >= deadline:
                            raise LibraryError('스위트가 다른 작업에서 사용 중입니다', 'SUITE_LOCKED', 409) from exc
                        time.sleep(0.05)
            else:
                import fcntl
                fcntl.flock(file, fcntl.LOCK_EX)
            _SUITE_HELD.keys = held | {key}
            try:
                from _tc_import_ops import recover_suite
                recover_suite(suite)
                yield
            finally:
                _SUITE_HELD.keys = held
                if os.name == 'nt':
                    file.seek(0)
                    msvcrt.locking(file.fileno(), msvcrt.LK_UNLCK, 1)
                else:
                    fcntl.flock(file, fcntl.LOCK_UN)


def without_auto(value):
    """Hide retired fields without rewriting legacy stored data."""
    if isinstance(value, dict):
        return {key: without_auto(item) for key, item in value.items() if key not in ('auto', 'default_auto')}
    if isinstance(value, list):
        return [without_auto(item) for item in value]
    return value


def _locked(function):
    @wraps(function)
    def guarded(suite, *args, **kwargs):
        with suite_lock(suite):
            return without_auto(function(suite, *args, **kwargs))
    return guarded


def _writes(function):
    """스위트가 있어야 하는 쓰기. update_state가 폴더를 만들어 주므로, 삭제 직후 도착한 요청
    (생성 작업 완료·다른 탭 편집)이 지운 스위트를 빈 폴더로 되살리지 않게 막는다.
    새 스위트를 만드는 것은 가져오기(import_cases·save_template)뿐이다."""
    @wraps(function)
    def guarded(suite, *args, **kwargs):
        with suite_lock(suite):
            if not suite_dir(suite).is_dir():
                raise LibraryError(f"스위트가 없습니다: {suite}", "SUITE_NOT_FOUND", 404)
            return without_auto(function(suite, *args, **kwargs))
    return guarded


def _cases_path(suite: str) -> Path:
    return suite_dir(suite) / "cases.json"


def _history_path(suite: str) -> Path:
    return suite_dir(suite) / "history.jsonl"


def list_suites() -> list[dict]:
    root = _paths.TC_LIBRARY_DIR
    if not root.exists():
        return []
    suites = []
    for d in sorted(p for p in root.iterdir() if p.is_dir() and not p.name.startswith("_")):
        with suite_lock(d.name):
            data = read_state(d / "cases.json")
            if data and not data.get("review_source_backfilled"):
                data = update_state(d / "cases.json", lambda cur: _backfill_review_source(d.name, cur))
        live = [c for c in data.get("cases", []) if not c.get("deleted")]
        suites.append({"suite": d.name, "sheets": data.get("sheets", []), "count": len(live),
                       "imported": sum(map(is_imported, live)),
                       "protected": d.name == DEFAULT_SUITE})
    return suites


@_locked
def load_cases(suite: str, *, include_deleted: bool = False) -> list[dict]:
    cases = read_state(_cases_path(suite)).get("cases", [])
    return cases if include_deleted else [c for c in cases if not c.get("deleted")]


@_locked
def load_profiles(suite: str) -> dict[str, TemplateProfile]:
    path = suite_dir(suite) / "template_profile.json"
    if not path.exists():
        return {}
    raw = json.loads(path.read_text(encoding="utf-8"))
    return {name: TemplateProfile.from_dict(p) for name, p in raw.items()}


@_locked
def get_case(suite: str, case_id: str) -> dict:
    for case in load_cases(suite, include_deleted=True):
        if case["case_id"] == case_id:
            return case
    raise LibraryError(f"케이스가 없습니다: {case_id}", "CASE_NOT_FOUND", 404)


# 누가 승인했는가: "import"(엑셀에서 가져와 자동 승인) / "human"(사람이 검토). 없으면 human으로 본다.
# 가져온 케이스는 검토 대상이 아니라 검토 상태를 바꿀 수 없다. 재가져오기는 이 값을 덮지 않는다.
REVIEW_IMPORT = "import"


def is_imported(case: dict) -> bool:
    return case.get("review_source") == REVIEW_IMPORT


def _backfill_review_source(suite: str, data: dict) -> dict:
    """review_source가 생기기 전 데이터를 변경 이력으로 한 번만 채운다. 가져오기로 만들어졌고
    사람이 검토 상태를 바꾼 기록이 없으면 import. 이력이 없으면 추측하지 않는다(human)."""
    created_by_import, status_edited = set(), set()
    path = _history_path(suite)
    for line in (path.read_text(encoding="utf-8").splitlines() if path.exists() else []):
        try:
            e = json.loads(line)
        except ValueError:
            continue
        if (e.get("kind") == "import" and e.get("before") is None) or \
                (e.get("kind") == "create" and e.get("after") == "import"):
            created_by_import.add(e["case_id"])
        elif e.get("field") == "status":
            status_edited.add(e["case_id"])
    for case in data.get("cases", []):
        if "review_source" not in case and case["case_id"] in created_by_import \
                and case["case_id"] not in status_edited and case["status"] == "approved":
            case["review_source"] = REVIEW_IMPORT
    return {**data, "review_source_backfilled": True}


@_writes
def set_review_source(suite: str, case_id: str, source: str) -> None:
    def mutate(data: dict) -> dict:
        _find(data, case_id)["review_source"] = source
        return data
    update_state(_cases_path(suite), mutate)


def _entry(case_id: str, field: str, before, after, actor: str, kind: str = "edit") -> dict:
    return {"history_id": "h_" + secrets.token_hex(6), "case_id": case_id, "field": field,
            "before": before, "after": after, "actor": actor, "kind": kind, "at": now_iso()}


@_writes
def _append_history(suite: str, entries: list[dict]) -> None:
    if not entries:
        return
    path = _history_path(suite)
    path.parent.mkdir(parents=True, exist_ok=True)
    with _paths._file_lock(path.with_suffix(".lock"), path):
        with path.open("a", encoding="utf-8") as fh:
            for entry in entries:
                fh.write(json.dumps(entry, ensure_ascii=False) + "\n")


def _apply(case: dict, changes: dict, actor: str) -> list[dict]:
    entries = []
    if "status" in changes and changes["status"] != case["status"] and is_imported(case):
        raise LibraryError("엑셀에서 가져온 케이스는 검토 상태를 바꿀 수 없습니다", "IMPORTED_REVIEW_LOCKED", 409)
    for field, value in changes.items():
        if field not in EDITABLE_FIELDS:
            raise LibraryError(f"수정할 수 없는 필드입니다: {field}", "FIELD_NOT_EDITABLE")
        if case.get(field) != value:
            entries.append(_entry(case["case_id"], field, case.get(field), value, actor))
            case[field] = value
    if entries:
        case["rev"] += 1
        case["updated_at"] = now_iso()
    return entries


def _find(data: dict, case_id: str, *, deleted: bool = False) -> dict:
    for case in data.get("cases", []):
        if case["case_id"] == case_id and bool(case.get("deleted")) == deleted:
            return case
    raise LibraryError(f"케이스가 없습니다: {case_id}", "CASE_NOT_FOUND", 404)


@_locked
def save_template(suite: str, xlsx_path: Path, profiles: dict[str, TemplateProfile]) -> None:
    target = suite_dir(suite)
    target.mkdir(parents=True, exist_ok=True)
    shutil.copy2(xlsx_path, target / "template.xlsx")
    (target / "template_profile.json").write_text(
        json.dumps({k: v.to_dict() for k, v in profiles.items()}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


@_locked
def import_cases(suite: str, sheets: list[str], incoming: list[dict], actor: str) -> dict:
    """case_id가 같으면 갱신, 없으면 추가 (F6.7). 같은 내용이면 건드리지 않는다."""
    summary = {"created": 0, "updated": 0, "unchanged": 0}
    history: list[dict] = []

    def mutate(data: dict) -> dict:
        data.setdefault("suite", suite)
        data["sheets"] = list(dict.fromkeys(data.get("sheets", []) + sheets))
        cases = data.setdefault("cases", [])
        by_id = {c["case_id"]: c for c in cases}
        for new in incoming:
            old = by_id.get(new["case_id"])
            if old is None:
                new = {**new, "review_source": REVIEW_IMPORT}
                cases.append(new)
                by_id[new["case_id"]] = new
                history.append(_entry(new["case_id"], "*", None, "import", actor, "create"))
                summary["created"] += 1
                continue
            changes = {f: new[f] for f in EDITABLE_FIELDS if f != "status" and f in new and old.get(f) != new[f]}
            if old.get("deleted"):
                old["deleted"] = False
                changes.setdefault("status", old["status"])
            if changes:
                history.extend(_apply(old, changes, actor))
                summary["updated"] += 1
            else:
                summary["unchanged"] += 1
        return data

    update_state(_cases_path(suite), mutate)
    _append_history(suite, history)
    return summary


@_writes
def patch_case(suite: str, case_id: str, base_rev: int, changes: dict, actor: str) -> dict:
    out: dict = {}

    def mutate(data: dict) -> dict:
        case = _find(data, case_id)
        if case["rev"] != base_rev:
            raise RevConflict(dict(case))
        out["entries"] = _apply(case, changes, actor)
        out["case"] = dict(case)
        return data

    update_state(_cases_path(suite), mutate)
    _append_history(suite, out["entries"])
    return out["case"]


@_writes
def bulk_patch(suite: str, items: list[dict], changes: dict, actor: str) -> dict:
    """rev가 맞는 케이스만 바꾸고 나머지는 conflicts로 돌려준다 (피드백 #15)."""
    result: dict = {"updated": [], "conflicts": [], "skipped": []}
    history: list[dict] = []

    def mutate(data: dict) -> dict:
        for item in items:
            case = _find(data, item["case_id"])
            if case["rev"] != item["rev"]:
                result["conflicts"].append(dict(case))
                continue
            try:
                history.extend(_apply(case, changes, actor))
            except LibraryError as exc:
                if exc.code != "IMPORTED_REVIEW_LOCKED":
                    raise
                result["skipped"].append(case["case_id"])     # 가져온 케이스는 검토 상태 일괄 변경에서 건너뜀
                continue
            result["updated"].append(case["case_id"])
        return data

    update_state(_cases_path(suite), mutate)
    _append_history(suite, history)
    return result


def _prefix_for(cases: list[dict], sheet: str) -> str:
    for case in cases:
        if case["sheet"] == sheet and "_" in case["case_id"]:
            return case["case_id"].rsplit("_", 1)[0]
    return "TC"


@_writes
def create_case(suite: str, fields: dict, actor: str, *, after: str | None = None) -> dict:
    from _tc_model import new_case

    out: dict = {}

    def mutate(data: dict) -> dict:
        cases = data.setdefault("cases", [])
        sheet = fields.get("sheet") or (data.get("sheets") or [""])[0]
        case_id = next_case_id(_prefix_for(cases, sheet), {c["case_id"] for c in cases})
        base = {k: v for k, v in fields.items()
                if k in EDITABLE_FIELDS or k in ("source_refs", "draft_meta")}
        base.update(case_id=case_id, sheet=sheet, status="draft")
        case = new_case(**base)
        index = next((i + 1 for i, c in enumerate(cases) if c["case_id"] == after), len(cases))
        cases.insert(index, case)
        out["case"] = dict(case)
        return data

    update_state(_cases_path(suite), mutate)
    _append_history(suite, [_entry(out["case"]["case_id"], "*", None, "create", actor, "create")])
    return out["case"]


@_writes
def duplicate_case(suite: str, case_id: str, actor: str) -> dict:
    source = get_case(suite, case_id)
    fields = {f: source[f] for f in EDITABLE_FIELDS if f in source}
    fields["source_tc_id"] = ""
    fields["source_refs"] = list(source.get("source_refs", []))
    created = create_case(suite, fields, actor, after=case_id)
    return created


@_writes
def delete_cases(suite: str, items: list[dict], actor: str) -> dict:
    result: dict = {"deleted": [], "conflicts": []}
    history: list[dict] = []

    def mutate(data: dict) -> dict:
        for item in items:
            case = _find(data, item["case_id"])
            if case["rev"] != item["rev"]:
                result["conflicts"].append(dict(case))
                continue
            case["deleted"] = True
            case["rev"] += 1
            case["updated_at"] = now_iso()
            history.append(_entry(case["case_id"], "deleted", False, True, actor, "delete"))
            result["deleted"].append(case["case_id"])
        return data

    update_state(_cases_path(suite), mutate)
    _append_history(suite, history)
    return result


@_writes
def restore_case(suite: str, case_id: str, actor: str) -> dict:
    out: dict = {}

    def mutate(data: dict) -> dict:
        case = _find(data, case_id, deleted=True)
        case["deleted"] = False
        case["rev"] += 1
        case["updated_at"] = now_iso()
        out["case"] = dict(case)
        return data

    update_state(_cases_path(suite), mutate)
    _append_history(suite, [_entry(case_id, "deleted", True, False, actor, "restore")])
    return out["case"]


@_locked
def history(suite: str, case_id: str) -> list[dict]:
    path = _history_path(suite)
    if not path.exists():
        return []
    entries = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line]
    return [e for e in reversed(entries) if e["case_id"] == case_id and e.get("field") not in ("auto", "default_auto")]


@_writes
def revert(suite: str, case_id: str, history_id: str, base_rev: int, actor: str) -> dict:
    entry = next((e for e in history(suite, case_id) if e["history_id"] == history_id), None)
    if entry is None or entry["kind"] != "edit":
        raise LibraryError("되돌릴 수 없는 이력입니다", "HISTORY_NOT_REVERTIBLE")
    return patch_case(suite, case_id, base_rev, {entry["field"]: entry["before"]}, actor)

def with_issues(case: dict) -> dict:
    from _tc_profiles import style_issues_for   # _tc_profiles가 이 모듈을 import한다
    issues = validate_case(case) + style_issues_for(case)
    return {**without_auto(case), "issues": issues, "has_error": any(i["level"] == "error" for i in issues)}


def filter_cases(cases: list[dict], query: dict) -> list[dict]:
    """GET /api/tc-library/{suite} 필터 (F5.4). query 값은 문자열."""
    q = (query.get("q") or "").strip()
    path = [p for p in (query.get("path") or "").split("/") if p]
    out = []
    for case in map(with_issues, cases):
        if query.get("sheet") and case["sheet"] != query["sheet"]:
            continue
        keys = [case["sheet"], *[p for p in case["path"] if p], case["feature"]]
        if path and keys[: len(path)] != path:
            continue
        status = query.get("status")
        if status == "imported" and not is_imported(case):
            continue
        if status and status != "imported" and (case["status"] != status or (status == "approved" and is_imported(case))):
            continue   # '승인' = 사람이 승인한 것만. 가져온 케이스는 '가져옴'으로 따로 거른다
        if "execution_result" in query and query["execution_result"] != "*" \
                and case["execution_result"] != query["execution_result"]:
            continue
        if query.get("priority") and (case["priority"] or "-") != query["priority"]:
            continue
        if query.get("source") and not any(r.startswith(query["source"] + ":")
                                           for r in case["source_refs"]):
            continue
        if query.get("invalid") == "1" and not case["has_error"]:
            continue
        if query.get("job") and case.get("draft_meta", {}).get("job_id") != query["job"]:
            continue
        if query.get("needs_review") == "1" and not case.get("flags", {}).get("source_change"):
            continue
        if q:
            haystack = json.dumps([case["feature"], case["precondition"], case["steps"],
                                   case["expected"], case["bullets"]], ensure_ascii=False)
            if q not in haystack:
                continue
        out.append(case)
    return out


def build_tree(cases: list[dict], branches: list[dict] | None = None) -> list[dict]:
    """시트 › 대분류 › 중분류 › 소분류 › 기능 트리 + 가지별 집계 (F2.2, F5.1)."""
    roots: list[dict] = []
    index: dict[tuple, dict] = {}
    for case in map(with_issues, cases):
        keys = [case["sheet"], *[p for p in case["path"] if p], case["feature"]]
        children = roots
        for depth, name in enumerate(keys):
            key = tuple(keys[: depth + 1])
            node = index.get(key)
            if node is None:
                level = "sheet" if depth == 0 else "feature" if depth == len(keys) - 1 else f"l{depth}"
                node = {"name": name, "level": level, "path": list(key), "count": 0,
                        "draft": 0, "needs_review": 0, "invalid": 0, "children": []}
                index[key] = node
                children.append(node)
            node["count"] += 1
            node["draft"] += case["status"] == "draft"
            node["needs_review"] += case["status"] == "needs_review" or bool(case.get("flags", {}).get("source_change"))
            node["invalid"] += case["has_error"]
            children = node["children"]
    for branch in branches or []:
        keys = [branch['sheet'], *[p for p in branch['path'] if p]]
        children = roots
        for depth, name in enumerate(keys):
            key = tuple(keys[:depth + 1])
            node = index.get(key)
            if node is None:
                node = {'name': name, 'level': 'sheet' if depth == 0 else f'l{depth}',
                        'path': list(key), 'count': 0, 'draft': 0, 'needs_review': 0,
                        'invalid': 0, 'children': []}
                index[key] = node
                children.append(node)
            elif node['level'] == 'feature':
                node['level'] = f'l{depth}'
            children = node['children']
    return roots


def _insert_index(cases: list[dict], sheet: str, path: list[str]) -> int:
    """같은 시트에서 경로가 가장 많이 겹치는 가지의 마지막 케이스 바로 뒤."""
    best, index = -1, len(cases)
    for i, case in enumerate(cases):
        if case["sheet"] != sheet or case.get("deleted"):
            continue
        depth = 0
        while depth < 3 and case["path"][depth] == path[depth] and path[depth]:
            depth += 1
        if depth >= best:
            best, index = depth, i + 1
    return index


@_writes
def add_drafts(suite: str, drafts: list[dict], actor: str) -> list[dict]:
    """생성 초안을 한 번에 추가한다 (Phase 2 G6). 대상 가지 끝에 순서대로 넣는다."""
    from _tc_model import new_case

    created: list[dict] = []

    def mutate(data: dict) -> dict:
        cases = data.setdefault("cases", [])
        ids = {c["case_id"] for c in cases}
        for draft in drafts:
            sheet = draft["sheet"]
            case_id = next_case_id(_prefix_for(cases, sheet), ids)
            ids.add(case_id)
            fields = {k: v for k, v in draft.items()
                      if k in EDITABLE_FIELDS or k in ("source_refs", "draft_meta")}
            fields.update(case_id=case_id, status="draft")
            case = new_case(**fields)
            cases.insert(_insert_index(cases, sheet, case["path"]), case)
            created.append(dict(case))
        return data

    update_state(_cases_path(suite), mutate)
    _append_history(suite, [_entry(c["case_id"], "*", None, "generate", actor, "create") for c in created])
    return created


@_writes
def set_draft_meta(suite: str, case_id: str, changes: dict) -> dict:
    """초안 메타만 바꾼다 (rev를 올리지 않는 내부용: 중복 표시 해제 등)."""
    out: dict = {}

    def mutate(data: dict) -> dict:
        case = _find(data, case_id)
        case["draft_meta"] = {**case.get("draft_meta", {}), **changes}
        out["case"] = dict(case)
        return data

    update_state(_cases_path(suite), mutate)
    return out["case"]


@_writes
def add_source_refs(suite: str, case_id: str, refs: list[str]) -> dict:
    """출처만 덧붙인다 (내용 변경이 아니므로 rev를 올리지 않는다)."""
    out: dict = {}

    def mutate(data: dict) -> dict:
        case = _find(data, case_id)
        case["source_refs"] = list(dict.fromkeys(case["source_refs"] + refs))
        out["case"] = dict(case)
        return data

    update_state(_cases_path(suite), mutate)
    return out["case"]


@_writes
def set_flag(suite: str, case_id: str, name: str, value) -> dict:
    """검토 표시를 켜고 끈다 (value=None이면 지움). 내용 변경이 아니므로 rev를 올리지 않는다."""
    out: dict = {}

    def mutate(data: dict) -> dict:
        case = _find(data, case_id)
        flags = dict(case.get("flags", {}))
        if value is None:
            flags.pop(name, None)
        else:
            flags[name] = value
        case["flags"] = flags
        out["case"] = dict(case)
        return data

    update_state(_cases_path(suite), mutate)
    return out["case"]


@_writes
def replace_source_version(suite: str, case_id: str, old_ref: str, new_ref: str) -> dict:
    """출처의 버전 부분만 바꾼다: "conf:1@v14#§2" → "conf:1@v15#§2" (섹션 앵커는 유지)."""
    out: dict = {}

    def mutate(data: dict) -> dict:
        case = _find(data, case_id)
        case["source_refs"] = [new_ref + r[len(old_ref):] if r.startswith(old_ref) else r
                               for r in case["source_refs"]]
        out["case"] = dict(case)
        return data

    update_state(_cases_path(suite), mutate)
    return out["case"]


@_writes
def rename_sheet(suite: str, old: str, name: str, actor: str) -> None:
    """Rename the worksheet and its cases without changing case IDs."""
    import re
    import tempfile
    import openpyxl
    from _tc_generate import active_job

    name = name.strip() if isinstance(name, str) else ''
    if not name or len(name) > 31 or re.search(r"[\\/*?:\[\]]", name) or name.startswith("'") or name.endswith("'"):
        raise LibraryError('시트 이름은 1~31자이며 \\ / * ? : [ ]를 포함할 수 없습니다.', 'INVALID_SHEET')
    running = active_job()
    if running and running['suite'] == suite:
        raise LibraryError('생성이 끝난 후 시트 이름을 변경하세요.', 'JOB_RUNNING', 409)
    root = suite_dir(suite)
    workbook = root / 'template.xlsx'
    profiles_path = root / 'template_profile.json'
    profiles = read_state(profiles_path)
    if old not in profiles:
        raise LibraryError('시트가 없습니다.', 'SHEET_NOT_FOUND', 404)
    if name == old:
        return
    wb = openpyxl.load_workbook(workbook)
    try:
        if any(title.casefold() == name.casefold() for title in wb.sheetnames if title != old):
            raise LibraryError('이미 존재하는 시트 이름입니다.', 'SHEET_EXISTS', 409)
        ws = wb[old]
        if old.casefold() == name.casefold():
            ws.title = '_rename_' + secrets.token_hex(4)
        ws.title = name
        with tempfile.TemporaryDirectory(prefix='tc-sheet-') as temp:
            renamed = Path(temp) / 'template.xlsx'
            wb.save(renamed)
            json_paths = [_cases_path(suite), profiles_path]
            md_path = root / 'md_export.json'
            if md_path.exists():
                json_paths.append(md_path)
            jobs_root = _paths.TC_LIBRARY_DIR / '_jobs'
            for job_path in jobs_root.glob('*/status.json'):
                job = read_state(job_path)
                if job.get('suite') == suite and job.get('target', {}).get('sheet') == old:
                    json_paths.append(job_path)
            originals = {path: read_state(path) for path in json_paths}
            original_workbook = workbook.read_bytes()
            history = []
            try:
                shutil.copy2(renamed, workbook)
                def mutate(data):
                    if name in data.get('sheets', []):
                        raise LibraryError('이미 존재하는 시트 이름입니다.', 'SHEET_EXISTS', 409)
                    data['sheets'] = [name if sheet == old else sheet for sheet in data.get('sheets', [])]
                    for branch in data.get('branches', []):
                        if branch['sheet'] == old:
                            branch['sheet'] = name
                    for case in data.get('cases', []):
                        if case['sheet'] == old:
                            history.extend(_apply(case, {'sheet': name}, actor))
                    return data
                new_profiles = {name if key == old else key: {**value, 'sheet': name if key == old else value['sheet']} for key, value in profiles.items()}
                update_state(profiles_path, lambda _: new_profiles)
                for path in json_paths[2:]:
                    def update_related(data):
                        for group in data.get('groups', []):
                            if group['path'] and group['path'][0] == old:
                                group['path'][0] = name
                        if data.get('target', {}).get('sheet') == old:
                            data['target']['sheet'] = name
                        return data
                    update_state(path, update_related)
                update_state(_cases_path(suite), mutate)
            except Exception:
                workbook.write_bytes(original_workbook)
                for path, original in originals.items():
                    update_state(path, lambda _, original=original: original)
                raise
            _append_history(suite, history)
    finally:
        wb.close()


@_locked
def load_branches(suite: str) -> list[dict]:
    return read_state(_cases_path(suite)).get('branches', [])


@_writes
def add_branch(suite: str, sheet: str, path: list[str]) -> dict:
    if sheet not in load_profiles(suite):
        raise LibraryError('시트를 선택하세요.', 'SHEET_NOT_FOUND', 404)
    if not isinstance(path, list) or not 1 <= len(path) <= 3 or any(not isinstance(p, str) for p in path):
        raise LibraryError('대·중·소분류 이름을 입력하세요.', 'INVALID_BRANCH')
    path = [p.strip() for p in path] + [''] * (3 - len(path))
    if not path[0] or any(path[i] and not path[i - 1] for i in range(1, 3)):
        raise LibraryError('상위 분류부터 차례로 입력하세요.', 'INVALID_BRANCH')
    if any(len(p) > 80 or '/' in p or any(ord(c) < 32 for c in p) for p in path):
        raise LibraryError('분류 이름은 80자 이내이며 / 또는 제어 문자를 포함할 수 없습니다.', 'INVALID_BRANCH')
    branch = {'sheet': sheet, 'path': path}
    def mutate(data):
        branches = data.setdefault('branches', [])
        if branch not in branches:
            branches.append(branch)
        return data
    update_state(_cases_path(suite), mutate)
    return branch


@_writes
def add_sheet(suite: str, name: str) -> None:
    """Add an empty worksheet using the existing template's formatting."""
    import re
    import tempfile
    from copy import deepcopy
    from dataclasses import replace
    import openpyxl
    from _tc_xlsx_export import write_sheet

    name = name.strip() if isinstance(name, str) else ''
    if not name or len(name) > 31 or re.search(r"[\\/*?:\[\]]", name) or name.startswith("'") or name.endswith("'"):
        raise LibraryError('시트 이름은 1~31자이며 \\ / * ? : [ ]를 포함할 수 없습니다.', 'INVALID_SHEET')
    root = suite_dir(suite)
    profiles = load_profiles(suite)
    if not profiles:
        raise LibraryError('먼저 빈 엑셀 양식을 가져오세요.', 'TEMPLATE_REQUIRED')
    workbook = root / 'template.xlsx'
    profile_path = root / 'template_profile.json'
    wb = openpyxl.load_workbook(workbook)
    try:
        if any(title.casefold() == name.casefold() for title in wb.sheetnames):
            raise LibraryError('이미 존재하는 시트 이름입니다.', 'SHEET_EXISTS', 409)
        source = next(iter(profiles.values()))
        ws = wb.copy_worksheet(wb[source.sheet])
        ws.title = name
        ws.data_validations = deepcopy(wb[source.sheet].data_validations)
        profile = replace(source, sheet=name)
        write_sheet(ws, profile, [])
        for col in range(1, ws.max_column + 1):
            ws.cell(profile.style_row, col)._style = deepcopy(wb[source.sheet].cell(source.style_row, col)._style)
        with tempfile.TemporaryDirectory(prefix='tc-add-sheet-') as temp:
            out = Path(temp) / 'template.xlsx'
            wb.save(out)
            previous = {path: read_state(path) for path in [profile_path, _cases_path(suite)]}
            original_workbook = workbook.read_bytes()
            try:
                shutil.copy2(out, workbook)
                update_state(profile_path, lambda data: {**data, name: profile.to_dict()})
                def mutate(data):
                    data['sheets'] = list(dict.fromkeys([*data.get('sheets', []), name]))
                    return data
                update_state(_cases_path(suite), mutate)
            except Exception:
                workbook.write_bytes(original_workbook)
                for path, original in previous.items():
                    update_state(path, lambda _, original=original: original)
                raise
    finally:
        wb.close()
