"""생성용 소스 번들 — 업로드 파일·붙여넣기를 markdown으로 정리 (PRD F1.1, F1.2, F1.5, F1.6).

state/tc_library/_sources/{bundle_id}/
  manifest.json   {"bundle_id", "created_at", "sources": [entry…], "removed": [entry…]}
  NN_{slug}.md    소스 1개 = 파일 1개 (정규화한 markdown)
entry: {source_id, kind, title, ref, version, sha256, chars, pages, truncated, warnings, file}
ref 형식: "file:{sha256 앞 12자}" (파일), "paste:{sha256 앞 12자}" (붙여넣기). 섹션은 ref + "#§N".
removed: 목록에서 뺀 소스. 이미 끝난 작업의 초안이 원문을 계속 참조하므로 파일은 지우지 않는다.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import secrets
import zipfile
from pathlib import Path

import _paths
from _state import update_state
from _tc_library import LibraryError
from _tc_model import now_iso

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_PASTE_BYTES = 1 * 1024 * 1024
MAX_PDF_PAGES = 200
MAX_DOCX_UNCOMPRESSED = 100 * 1024 * 1024
MAX_SOURCE_CHARS = 400_000            # 소스 1개당 보관 상한. 넘으면 잘라내고 truncated 표시
ALLOWED_EXT = (".pdf", ".docx", ".md", ".txt")
_HEADING = re.compile(r"^(#{1,3})\s+(.+?)\s*$")


class SourceError(LibraryError):
    pass


def sources_root() -> Path:
    return _paths.TC_LIBRARY_DIR / "_sources"


def _bundle_dir(bundle_id: str) -> Path:
    if not re.fullmatch(r"src_[0-9a-f]{12}", bundle_id or ""):
        raise SourceError("소스 묶음 id가 올바르지 않습니다", "INVALID_BUNDLE")
    return sources_root() / bundle_id


def new_bundle() -> str:
    bundle_id = "src_" + secrets.token_hex(6)
    d = _bundle_dir(bundle_id)
    d.mkdir(parents=True, exist_ok=True)
    update_state(d / "manifest.json",
                 lambda _: {"bundle_id": bundle_id, "created_at": now_iso(), "sources": []})
    return bundle_id


def load_bundle(bundle_id: str) -> dict:
    path = _bundle_dir(bundle_id) / "manifest.json"
    if not path.exists():
        raise SourceError("소스 묶음이 없습니다", "BUNDLE_NOT_FOUND", 404)
    return json.loads(path.read_text(encoding="utf-8"))


def all_entries(manifest: dict) -> list[dict]:
    """목록에 있는 소스 + 목록에서 뺀 소스 (원문 조회용)."""
    return manifest["sources"] + manifest.get("removed", [])


def read_text(bundle_id: str, source_id: str) -> str:
    entry = next((s for s in all_entries(load_bundle(bundle_id)) if s["source_id"] == source_id), None)
    if entry is None:
        raise SourceError("소스가 없습니다", "SOURCE_NOT_FOUND", 404)
    return (_bundle_dir(bundle_id) / entry["file"]).read_text(encoding="utf-8")


# ── 추출 ────────────────────────────────────────────────────────
def _pdf_to_markdown(data: bytes) -> tuple[str, int, list[str]]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if len(reader.pages) > MAX_PDF_PAGES:
        raise SourceError(f"PDF가 {MAX_PDF_PAGES}쪽을 넘습니다 ({len(reader.pages)}쪽)", "PDF_TOO_LONG", 413)
    parts, empty = [], 0
    for number, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if not text:
            empty += 1
            continue
        parts.append(f"# {number}쪽\n\n{text}")
    warnings = [f"텍스트 없는 쪽 {empty}개 (스캔 PDF OCR 미지원)"] if empty else []
    return "\n\n".join(parts), len(reader.pages), warnings


def _docx_to_markdown(data: bytes) -> tuple[str, int, list[str]]:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise SourceError("DOCX 형식이 아닙니다", "UNSUPPORTED_FILE") from exc
    if "word/document.xml" not in archive.namelist():
        raise SourceError("DOCX 형식이 아닙니다", "UNSUPPORTED_FILE")
    if sum(i.file_size for i in archive.infolist()) > MAX_DOCX_UNCOMPRESSED:
        raise SourceError("압축을 푼 크기가 너무 큽니다", "DOCX_TOO_LARGE", 413)
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines: list[str] = []
    body = doc.element.body
    tables = iter(doc.tables)
    paragraphs = iter(doc.paragraphs)
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = next(paragraphs)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            level = re.search(r"(\d)$", style) if style.lower().startswith(("heading", "제목")) else None
            lines.append(f"{'#' * min(int(level.group(1)), 3)} {text}" if level else text)
        elif tag == "tbl":
            table = next(tables)
            rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in table.rows]
            if rows:
                lines.append("| " + " | ".join(rows[0]) + " |")
                lines.append("|" + "---|" * len(rows[0]))
                lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n\n".join(lines), 0, []


def extract(filename: str, data: bytes) -> tuple[str, int, list[str]]:
    """(markdown, 쪽 수, 경고). 확장자와 매직바이트가 모두 맞아야 한다."""
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXT:
        raise SourceError(f"지원하지 않는 형식입니다: {ext or filename}", "UNSUPPORTED_FILE")
    if ext == ".pdf":
        if not data.startswith(b"%PDF"):
            raise SourceError("PDF 형식이 아닙니다", "UNSUPPORTED_FILE")
        return _pdf_to_markdown(data)
    if ext == ".docx":
        if not data.startswith(b"PK"):
            raise SourceError("DOCX 형식이 아닙니다", "UNSUPPORTED_FILE")
        return _docx_to_markdown(data)
    try:
        return data.decode("utf-8-sig"), 0, []
    except UnicodeDecodeError as exc:
        raise SourceError("UTF-8 텍스트가 아닙니다", "UNSUPPORTED_FILE") from exc


def split_sections(markdown: str) -> list[dict]:
    """제목(#~###) 기준 섹션. 제목이 없으면 문서 전체가 섹션 1개. anchor는 "§1", "§2"…"""
    sections: list[dict] = []
    title, buf = "", []

    def flush() -> None:
        text = "\n".join(buf).strip()
        if text or title:
            sections.append({"anchor": f"§{len(sections) + 1}", "title": title or "본문", "text": text})

    for line in markdown.splitlines():
        m = _HEADING.match(line)
        if m:
            flush()
            title, buf = m.group(2), []
        else:
            buf.append(line)
    flush()
    return sections


# ── 추가·삭제 ───────────────────────────────────────────────────
def _slug(text: str) -> str:
    return re.sub(r"[^\w가-힣]+", "_", text).strip("_")[:40] or "source"


def _add(bundle_id: str, kind: str, title: str, markdown: str, digest: str,
         pages: int, warnings: list[str], *, ref: str = "", version: str = "",
         origin: str = "", assets: dict[str, bytes] | None = None) -> dict:
    d = _bundle_dir(bundle_id)
    if not (d / "manifest.json").exists():
        raise SourceError("소스 묶음이 없습니다", "BUNDLE_NOT_FOUND", 404)
    truncated = len(markdown) > MAX_SOURCE_CHARS
    if truncated:
        markdown = markdown[:MAX_SOURCE_CHARS]
        warnings = warnings + [f"{MAX_SOURCE_CHARS:,}자에서 잘랐습니다"]
    entry: dict = {}

    def mutate(manifest: dict) -> dict:
        source_ref = ref or f"{kind}:{digest[:12]}"
        if any(s["ref"] == source_ref for s in manifest["sources"]):
            raise SourceError("이미 추가한 소스입니다", "SOURCE_EXISTS", 409)
        # 지운 번호를 다시 쓰지 않는다 (지운 뒤 추가해도 s01·s02가 겹치지 않게)
        n = max([int(x["source_id"][1:]) for x in manifest["sources"]] + [manifest.get("last_n", 0)]) + 1
        manifest["last_n"] = n
        filename = f"{n:02d}_{_slug(title)}.md"
        (d / filename).write_text(markdown, encoding="utf-8")
        asset_names = []
        for name, data in (assets or {}).items():
            safe = f"{n:02d}_{_slug(name)}.png"
            (d / "assets").mkdir(exist_ok=True)
            (d / "assets" / safe).write_bytes(data)
            asset_names.append(safe)
        entry.update({
            "source_id": f"s{n:02d}", "kind": kind, "title": title, "ref": source_ref,
            "version": version or digest[:12], "sha256": digest, "chars": len(markdown), "pages": pages,
            "sections": len(split_sections(markdown)), "truncated": truncated,
            "warnings": warnings, "file": filename, "origin": origin, "assets": asset_names,
            "added_at": now_iso(),
        })
        manifest["sources"].append(dict(entry))
        manifest["removed"] = [s for s in manifest.get("removed", []) if s["ref"] != source_ref]
        return manifest

    update_state(d / "manifest.json", mutate)
    return entry


def add_file(bundle_id: str, filename: str, data: bytes) -> dict:
    if len(data) > MAX_FILE_BYTES:
        raise SourceError("25MB를 넘습니다", "PAYLOAD_TOO_LARGE", 413)
    name = Path(filename).name
    markdown, pages, warnings = extract(name, data)
    if not markdown.strip():
        raise SourceError("추출한 텍스트가 없습니다", "EMPTY_SOURCE")
    return _add(bundle_id, "file", name, markdown, hashlib.sha256(data).hexdigest(), pages, warnings)


def add_paste(bundle_id: str, text: str) -> dict:
    raw = (text or "").encode("utf-8")
    if len(raw) > MAX_PASTE_BYTES:
        raise SourceError("붙여넣기는 1MB까지입니다", "PAYLOAD_TOO_LARGE", 413)
    if not text.strip():
        raise SourceError("붙여넣은 내용이 없습니다", "EMPTY_SOURCE")
    title = text.strip().splitlines()[0][:30]
    return _add(bundle_id, "paste", title, text, hashlib.sha256(raw).hexdigest(), 0, [])


def add_fetched(bundle_id: str, *, kind: str, title: str, markdown: str, ref: str, version: str,
                origin: str, warnings: list[str] | None = None, pages: int = 0,
                assets: dict[str, bytes] | None = None) -> dict:
    """원격 소스(url·conf·figma) 추가 (Phase 3). ref에는 버전이 들어간다: "conf:48213377@v14"."""
    if not markdown.strip():
        raise SourceError("가져온 문서에 본문이 없습니다", "EMPTY_SOURCE")
    digest = hashlib.sha256(markdown.encode("utf-8")).hexdigest()
    return _add(bundle_id, kind, title, markdown, digest, pages, list(warnings or []),
                ref=ref, version=version, origin=origin, assets=assets)


def asset_path(bundle_id: str, name: str) -> Path:
    if not re.fullmatch(r"\d\d_[\w가-힣]+\.png", name):
        raise SourceError("파일 이름이 올바르지 않습니다", "INVALID_ASSET")
    path = _bundle_dir(bundle_id) / "assets" / name
    if not path.exists():
        raise SourceError("이미지가 없습니다", "ASSET_NOT_FOUND", 404)
    return path


def find_source(ref: str) -> tuple[str, dict] | None:
    """버전 포함 ref("conf:123@v14")를 가진 소스를 모든 묶음에서 찾는다 (가장 최근 묶음 우선)."""
    root = sources_root()
    if not root.exists():
        return None
    for d in sorted(root.iterdir(), key=lambda p: p.stat().st_mtime, reverse=True):
        manifest = d / "manifest.json"
        if not manifest.exists():
            continue
        for entry in all_entries(json.loads(manifest.read_text(encoding="utf-8"))):
            if entry["ref"] == ref:
                return d.name, entry
    return None


def remove_source(bundle_id: str, source_id: str) -> None:
    """목록에서만 뺀다. 파일은 남겨 이미 만든 초안의 원문 발췌·재생성이 계속 동작하게 한다."""
    d = _bundle_dir(bundle_id)

    def mutate(manifest: dict) -> dict:
        gone = [s for s in manifest["sources"] if s["source_id"] == source_id]
        if not gone:
            raise SourceError("소스가 없습니다", "SOURCE_NOT_FOUND", 404)
        manifest["sources"] = [s for s in manifest["sources"] if s["source_id"] != source_id]
        manifest["removed"] = manifest.get("removed", []) + [{**gone[0], "removed_at": now_iso()}]
        return manifest

    update_state(d / "manifest.json", mutate)


def excerpt(bundle_id: str, ref: str) -> dict:
    """ref("file:abc#§2" 또는 "file:abc") → 해당 섹션 본문."""
    base, _, anchor = ref.partition("#")
    manifest = load_bundle(bundle_id)
    entry = next((s for s in all_entries(manifest) if s["ref"] == base), None)
    if entry is None:
        raise SourceError("출처를 이 소스 묶음에서 찾을 수 없습니다", "SOURCE_NOT_FOUND", 404)
    sections = split_sections(read_text(bundle_id, entry["source_id"]))
    section = next((s for s in sections if s["anchor"] == anchor), sections[0])
    return {"ref": ref, "title": entry["title"], "section": section["title"],
            "anchor": section["anchor"], "markdown": section["text"]}
