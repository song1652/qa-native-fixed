"""엑셀 TC 템플릿 분석 — 헤더 위치·컬럼·드롭다운·No. 수식 (PRD F2.1~F2.3, F2.6)."""
from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from pathlib import Path

HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "no": ("no.", "no"),
    "l1": ("대분류",),
    "l2": ("중분류",),
    "l3": ("소분류",),
    "feature": ("기능",),
    "precondition": ("사전 조건", "사전조건"),
    "steps": ("test step", "test steps", "테스트 절차"),
    "expected": ("expected result", "기대결과", "기대 결과"),
    "priority": ("우선순위",),
    "auto": ("auto",),
    "env": ("환경",),
    "note": ("기타", "비고"),
}
REQUIRED_COLUMNS = ("l1", "feature", "steps", "expected")


def _norm(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


@dataclass
class TemplateProfile:
    sheet: str
    header_row: int
    data_start_row: int
    columns: dict[str, int]                 # field → 1-based column
    result_columns: dict[str, int] = field(default_factory=dict)   # 플랫폼 → column
    validations: dict[str, list[str]] = field(default_factory=dict)  # field → 목록 값
    no_formula: str | None = None            # "=IF(H{r}<>\"\",ROW(B{r})-12, \"\")"
    style_row: int = 0
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, data: dict) -> "TemplateProfile":
        return cls(**data)


def _match_field(value) -> str | None:
    text = _norm(value)
    for key, aliases in HEADER_ALIASES.items():
        if text in aliases:
            return key
    return None


def detect_header_row(ws, max_scan: int = 40) -> int | None:
    """필수 헤더 4개(대분류·기능·Test Step·Expected Result)가 모두 있는 첫 행.

    요약 표(구분/COUNT/Pass…)는 필수 헤더가 없어서 자연히 건너뛴다 (명세 피드백 #4).
    """
    for row in ws.iter_rows(min_row=1, max_row=max_scan):
        found = {_match_field(c.value) for c in row}
        if all(key in found for key in REQUIRED_COLUMNS):
            return row[0].row
    return None


def analyze_sheet(ws) -> TemplateProfile | None:
    header_row = detect_header_row(ws)
    if header_row is None:
        return None
    columns: dict[str, int] = {}
    for cell in ws[header_row]:
        key = _match_field(cell.value)
        if key and key not in columns:
            columns[key] = cell.column

    result_columns: dict[str, int] = {}
    sub_row = header_row + 1
    if "env" in columns:
        col = columns["env"]
        while col <= ws.max_column:
            label = ws.cell(sub_row, col).value
            header_here = ws.cell(header_row, col).value
            if not label or (col != columns["env"] and header_here):
                break
            result_columns[str(label).strip()] = col
            col += 1
    data_start = sub_row + 1 if result_columns else sub_row

    validations: dict[str, list[str]] = {}
    col_to_field = {c: f for f, c in columns.items()}
    for platform_col in result_columns.values():
        col_to_field[platform_col] = "execution_result"
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list" or not dv.formula1:
            continue
        first = min(r.min_col for r in dv.sqref.ranges)
        key = col_to_field.get(first)
        if key:
            validations[key] = [v.strip() for v in dv.formula1.strip('"').split(",")]

    no_formula = None
    if "no" in columns:
        value = ws.cell(data_start, columns["no"]).value
        if isinstance(value, str) and value.startswith("="):
            no_formula = re.sub(rf"(?<=[A-Z]){data_start}(?!\d)", "{r}", value)

    warnings = []
    if getattr(ws, "_images", None):
        warnings.append(f"이미지 {len(ws._images)}개는 다시 저장하면 사라질 수 있습니다")
    if getattr(ws, "_charts", None):
        warnings.append(f"차트 {len(ws._charts)}개는 다시 저장하면 사라질 수 있습니다")
    if "priority" in validations and validations["priority"] != ["P0", "P1", "P2", "P3"]:
        warnings.append("우선순위 드롭다운을 내보낼 때 P0~P3으로 넓힙니다")

    return TemplateProfile(
        sheet=ws.title, header_row=header_row, data_start_row=data_start,
        columns=columns, result_columns=result_columns, validations=validations,
        no_formula=no_formula, style_row=data_start, warnings=warnings,
    )


def analyze_workbook(path: Path) -> dict[str, TemplateProfile]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path))
    try:
        profiles = {}
        for ws in wb.worksheets:
            profile = analyze_sheet(ws)
            if profile:
                profiles[ws.title] = profile
        return profiles
    finally:
        wb.close()
