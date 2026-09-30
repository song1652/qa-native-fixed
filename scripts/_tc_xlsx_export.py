"""라이브러리 케이스 → 템플릿 사본 xlsx (PRD F6.2~F6.6)."""
from __future__ import annotations

import re
from copy import copy
from datetime import date
from pathlib import Path

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import MultiCellRange

from _tc_model import PRIORITIES, RESULT_TO_EXCEL, format_steps, join_expected
from _tc_template import TemplateProfile

_HIER = ("l1", "l2", "l3", "feature")


def note_cell(case: dict) -> str:
    system = " | ".join(
        [f"id:{case['case_id']}"]
        + [f"src:{r}" for r in case.get("source_refs", []) if not r.startswith("xlsx:")]
    )
    return f"{case['note']}\n{system}" if case.get("note") else system


def _hier_key(case: dict, depth: int) -> tuple:
    keys = list(case["path"]) + [case["feature"]]
    return tuple(keys[: depth + 1])


def write_sheet(ws, profile: TemplateProfile, cases: list[dict]) -> int:
    """데이터 영역을 cases로 다시 쓰고 마지막 데이터 행 번호를 돌려준다."""
    start, cols = profile.data_start_row, profile.columns
    max_col = ws.max_column
    styles = [copy(ws.cell(profile.style_row, c)._style) for c in range(1, max_col + 1)]
    for rng in list(ws.merged_cells.ranges):
        if rng.max_row >= start:
            ws.unmerge_cells(str(rng))
    if ws.max_row >= start:
        ws.delete_rows(start, ws.max_row - start + 1)

    for i, case in enumerate(cases):
        r = start + i
        for c in range(1, max_col + 1):
            ws.cell(r, c)._style = copy(styles[c - 1])
        values = {
            "l1": case["path"][0], "l2": case["path"][1], "l3": case["path"][2],
            "feature": case["feature"], "precondition": case["precondition"],
            "steps": format_steps(case["steps"]),
            "expected": join_expected(case["expected"], case["bullets"]),
            "priority": case["priority"], "auto": case["auto"], "note": note_cell(case),
        }
        for key, val in values.items():
            if key in cols:
                ws.cell(r, cols[key]).value = val or None
        if profile.no_formula and "no" in cols:
            ws.cell(r, cols["no"]).value = profile.no_formula.format(r=r)
        excel_result = RESULT_TO_EXCEL.get(case["execution_result"])
        for col in profile.result_columns.values():
            ws.cell(r, col).value = excel_result

    last = start + max(len(cases), 1) - 1
    _merge_hierarchy(ws, profile, cases)
    _extend_validations(ws, profile, last)
    _rewrite_summary(ws, profile, last)
    return last


def _merge_hierarchy(ws, profile: TemplateProfile, cases: list[dict]) -> None:
    start = profile.data_start_row
    for depth, key in enumerate(_HIER):
        col = profile.columns.get(key)
        if col is None:
            continue
        run_start = 0
        for i in range(1, len(cases) + 1):
            same = (
                i < len(cases)
                and _hier_key(cases[i], depth) == _hier_key(cases[run_start], depth)
            )
            if same:
                continue
            value = (list(cases[run_start]["path"]) + [cases[run_start]["feature"]])[depth]
            if i - run_start > 1 and value:
                ws.merge_cells(
                    start_row=start + run_start, start_column=col,
                    end_row=start + i - 1, end_column=col,
                )
            run_start = i


def _extend_validations(ws, profile: TemplateProfile, last: int) -> None:
    start = profile.data_start_row
    priority_col = profile.columns.get("priority")
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list":
            continue
        lo = min(r.min_col for r in dv.sqref.ranges)
        hi = max(r.max_col for r in dv.sqref.ranges)
        dv.sqref = MultiCellRange(
            f"{get_column_letter(lo)}{start}:{get_column_letter(hi)}{last}"
        )
        if lo == priority_col:
            dv.formula1 = '"' + ",".join(PRIORITIES) + '"'


def _rewrite_summary(ws, profile: TemplateProfile, last: int) -> None:
    """요약 표 수식의 고정 범위·#REF!를 실제 데이터 범위로 다시 쓴다 (F6.3)."""
    start = profile.data_start_row
    platform_of_col: dict[int, str] = {}
    for row in ws.iter_rows(min_row=1, max_row=profile.header_row - 1):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() in profile.result_columns:
                platform_of_col.setdefault(cell.column, cell.value.strip())
    for row in ws.iter_rows(min_row=1, max_row=profile.header_row - 1):
        for cell in row:
            v = cell.value
            if not (isinstance(v, str) and v.startswith("=")):
                continue
            v = re.sub(r"\$A\$\d+:\$A\$\d+", f"$A${start}:$A${last}", v)
            if "#REF!" in v and cell.column in platform_of_col:
                col = get_column_letter(profile.result_columns[platform_of_col[cell.column]])
                v = v.replace("#REF!", f"${col}${start}:${col}${last}")
            cell.value = v


def append_history(wb, note: str, today: date) -> None:
    if "History" not in wb.sheetnames or not note.strip():
        return
    ws = wb["History"]
    last = max((c.row for c in ws["B"] if c.value not in (None, "")), default=2)
    for col in (2, 3):
        ws.cell(last + 1, col)._style = copy(ws.cell(last, col)._style)
    ws.cell(last + 1, 2).value = today.strftime("%y.%m.%d")
    ws.cell(last + 1, 3).value = note.strip()


def export_workbook(
    template_path: Path, profiles: dict[str, TemplateProfile],
    cases_by_sheet: dict[str, list[dict]], out_path: Path, *,
    history_note: str = "", today: date | None = None, drop_sheets: tuple[str, ...] = (),
) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(str(template_path))
    try:
        for name in drop_sheets:
            if name in wb.sheetnames:
                del wb[name]
        for sheet, profile in profiles.items():
            if sheet in wb.sheetnames and sheet not in drop_sheets:
                write_sheet(wb[sheet], profile, cases_by_sheet.get(sheet, []))
        append_history(wb, history_note, today or date.today())
        out_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(out_path))
    finally:
        wb.close()

def verify_export(
    out_path: Path, profiles: dict[str, TemplateProfile],
    cases_by_sheet: dict[str, list[dict]],
) -> list[dict]:
    """저장한 파일을 다시 열어 라이브러리와 같은지 검사한다 (F6.6)."""
    import openpyxl
    from _tc_xlsx_import import read_sheet_cases

    checks: list[dict] = []
    wb = openpyxl.load_workbook(str(out_path))
    try:
        for sheet, expected in cases_by_sheet.items():
            if sheet not in wb.sheetnames:
                continue
            ws, profile = wb[sheet], profiles[sheet]
            got = read_sheet_cases(ws, profile, source_name=out_path.name, prefix="X",
                                   existing_ids=set())
            key = lambda c: (c["case_id"], tuple(c["path"]), c["feature"], tuple(c["steps"]),
                             c["expected"], c["priority"], c["execution_result"], c["note"])
            same = [key(c) for c in got] == [key(c) for c in expected]
            checks.append({
                "level": "ok" if same else "error", "code": "ROUNDTRIP",
                "message": f"{sheet}: {len(got)}건 {'라이브러리와 일치' if same else '라이브러리와 다름'}",
            })
            refs = [
                c.coordinate for row in ws.iter_rows(max_row=profile.header_row)
                for c in row if isinstance(c.value, str) and "#REF!" in c.value
            ]
            checks.append({
                "level": "error" if refs else "ok", "code": "SUMMARY_REF",
                "message": f"{sheet}: 요약 수식 #REF! {len(refs)}개",
            })
    finally:
        wb.close()
    return checks
