"""Excel → library plans and recoverable multi-file transactions.

md commits continue to use _import_commit. These snapshots belong to library data.
"""
from __future__ import annotations

import base64
import copy
import hashlib
import io
import json
import os
import re
import secrets
from datetime import datetime
from pathlib import Path

import _paths
from _state import read_state, update_state
from _tc_library import LibraryError, _append_history, _entry, load_cases, suite_dir, suite_lock, without_auto
from _tc_model import EDITABLE_FIELDS, next_case_id, now_iso
from _tc_template import analyze_with_mapping, analyze_workbook, profile_from_mapping
from _tc_xlsx_import import import_workbook

FILES = ('cases.json', 'template.xlsx', 'template_profile.json', 'md_export.json', 'branches.json')


def _root() -> Path:
    return _paths.TC_LIBRARY_DIR / '_import_runs'


def _run_path(run_id: str) -> Path:
    if not re.fullmatch(r'libimp_[0-9a-f]{16}', str(run_id)):
        raise LibraryError('가져오기 작업 ID가 올바르지 않습니다', 'INVALID_RUN')
    return _root() / (run_id + '.json')


def _save(run: dict) -> None:
    update_state(_run_path(run['run_id']), lambda _: run)


def _read_run(path: Path) -> dict:
    if not path.exists():
        return {}
    try:
        run = json.loads(path.read_text(encoding='utf-8'))
        if not isinstance(run, dict) or not all(key in run for key in ('run_id', 'suite', 'status', 'created_at')):
            raise ValueError('required run fields are missing')
        return run
    except (ValueError, OSError) as exc:
        raise LibraryError('가져오기 이력이 손상되었습니다', 'RUN_CORRUPT', 409) from exc


def get_run(run_id: str) -> dict:
    run = _read_run(_run_path(run_id))
    if not run:
        raise LibraryError('가져오기 작업이 없습니다', 'RUN_NOT_FOUND', 404)
    with suite_lock(run['suite']):
        return _read_run(_run_path(run_id))


def public_run(run: dict, *, summary_only: bool = False) -> dict:
    hidden = {'before_snapshot', 'journal', 'before_hash', 'after_hash'}
    if summary_only:
        hidden.add('rows')
    result = {key: value for key, value in run.items() if key not in hidden}
    result['sources'] = [{key: value for key, value in source.items() if key not in ('profiles', 'meta_sha256', 'sha256')}
                         for source in run.get('sources', [])]
    return without_auto(result)


def list_runs(errors: list | None = None) -> list[dict]:
    runs = []
    for path in _root().glob('libimp_*.json'):
        try:
            runs.append(get_run(path.stem))
        except (LibraryError, ValueError, KeyError, OSError) as exc:
            if errors is not None:
                errors.append({'run_id': path.stem, 'code': getattr(exc, 'code', 'RUN_CORRUPT'), 'error': str(exc)})
    return sorted(runs, key=lambda run: (run['created_at'], run['run_id']), reverse=True)


def _snapshot(suite: str) -> dict:
    root = suite_dir(suite)
    snapshot = {}
    for name in FILES:
        path = root / name
        if not path.exists():
            snapshot[name] = None
            continue
        data = path.read_bytes()
        if name.endswith('.json'):
            try:
                data = json.dumps(json.loads(data), ensure_ascii=False, indent=2).encode()
            except ValueError as exc:
                raise LibraryError('스위트 데이터가 손상되었습니다', 'SUITE_CORRUPT', 409) from exc
        snapshot[name] = base64.b64encode(data).decode()
    return snapshot


def _fingerprint(snapshot: dict) -> str:
    return hashlib.sha256(json.dumps(snapshot, sort_keys=True).encode()).hexdigest()


def _bytes_write(path: Path, data: bytes | None) -> None:
    if data is None:
        path.unlink(missing_ok=True)
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + '.import-' + secrets.token_hex(4))
    try:
        with temporary.open('wb') as fh:
            fh.write(data)
            fh.flush()
            os.fsync(fh.fileno())
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _restore(suite: str, snapshot: dict) -> None:
    for name, encoded in snapshot.items():
        path = suite_dir(suite) / name
        if encoded is not None and name.endswith('.json'):
            data = json.loads(base64.b64decode(encoded))
            update_state(path, lambda _, data=data: data)
        else:
            _bytes_write(path, base64.b64decode(encoded) if encoded is not None else None)


def _history_once(suite: str, entries: list[dict]) -> None:
    path = suite_dir(suite) / 'history.jsonl'
    ids = {json.loads(line)['history_id'] for line in path.read_text().splitlines() if line} if path.exists() else set()
    _append_history(suite, [entry for entry in entries if entry['history_id'] not in ids])


def recover_suite(suite: str) -> None:
    """Called under the suite lock before every library read/write. Finish durable intent."""
    for path in _root().glob('libimp_*.json'):
        try:
            run = _read_run(path)
        except LibraryError:
            continue
        if run.get('suite') != suite or run.get('status') not in ('committing', 'rolling_back'):
            continue
        journal = run['journal']
        current = _snapshot(suite)
        if any(current[name] not in (journal['before'][name], journal['after'][name]) for name in FILES):
            raise LibraryError('중단된 가져오기 복구 중 변경을 발견했습니다', 'RECOVERY_CONFLICT', 409)
        _restore(suite, journal['after'])
        _history_once(suite, journal['history'])
        run['status'] = journal['final_status']
        run['after_hash'] = _fingerprint(_snapshot(suite))
        run.pop('journal', None)
        _save(run)


def _transaction(run: dict, after: dict, entries: list[dict], final_status: str) -> None:
    before = _snapshot(run['suite'])
    run['journal'] = {'before': before, 'after': after, 'history': entries, 'final_status': final_status}
    run['status'] = 'committing' if final_status == 'committed' else 'rolling_back'
    _save(run)
    # Roll forward after a process interruption; ordinary exceptions restore the pre-operation state.
    try:
        _restore(run['suite'], after)
    except Exception:
        _restore(run['suite'], before)
        run['status'] = 'preview_ready' if final_status == 'committed' else 'committed'
        run.pop('journal', None)
        _save(run)
        raise
    # A history append failure leaves the durable journal intact for idempotent recovery.
    _history_once(run['suite'], entries)
    run['status'] = final_status
    run['after_hash'] = _fingerprint(_snapshot(run['suite']))
    run.pop('journal', None)
    _save(run)


def _sources(body: dict) -> list[dict]:
    sources = body.get('sources') or [body]
    if not isinstance(sources, list) or not sources:
        raise LibraryError('파일을 하나 이상 선택하세요', 'INVALID_SOURCES')
    return sources


def _read_source(source: dict, body: dict) -> tuple[dict, list[dict], dict]:
    if not isinstance(source, dict):
        raise LibraryError('파일 설정이 올바르지 않습니다', 'INVALID_SOURCES')
    preview_id = str(source.get('preview_id', ''))
    if not re.fullmatch(r'imp_[0-9a-f]{12}', preview_id):
        raise LibraryError('미리보기 ID가 올바르지 않습니다', 'INVALID_PREVIEW')
    upload = _paths.TC_LIBRARY_DIR / '_uploads' / (preview_id + '.xlsx')
    if not upload.exists() or not upload.with_suffix('.json').exists():
        raise LibraryError('미리보기가 만료됐습니다', 'PREVIEW_EXPIRED', 410)
    meta = json.loads(upload.with_suffix('.json').read_text())
    filename = Path(meta['filename']).name
    profiles = analyze_with_mapping(upload, meta['mapping']) if meta.get('mapping') else analyze_workbook(upload)
    mappings = source.get('sheet_mappings') or body.get('sheet_mappings') or {}
    if not isinstance(mappings, dict):
        raise LibraryError('시트별 매핑이 올바르지 않습니다', 'INVALID_MAPPING')
    if mappings:
        import openpyxl
        wb = openpyxl.load_workbook(upload)
        try:
            for sheet, mapping in mappings.items():
                if not isinstance(mapping, dict):
                    raise LibraryError('시트 매핑이 올바르지 않습니다', 'INVALID_MAPPING')
                if sheet not in wb.sheetnames:
                    raise LibraryError('매핑할 시트가 없습니다', 'SHEET_NOT_FOUND')
                profiles[sheet] = profile_from_mapping(wb[sheet], mapping)
        finally:
            wb.close()
    sheets = source.get('sheets', [])
    if not isinstance(sheets, list) or not sheets or any(not isinstance(sheet, str) or sheet not in profiles for sheet in sheets):
        raise LibraryError('가져올 시트를 선택하세요', 'NO_SHEETS')
    prefixes = source.get('prefixes') or body.get('prefixes') or {}
    if not isinstance(prefixes, dict):
        raise LibraryError('접두어 설정이 올바르지 않습니다', 'INVALID_PREFIX')
    if any(not re.fullmatch(r'[A-Z][A-Z0-9]{0,7}', str(prefix)) for prefix in prefixes.values()):
        raise LibraryError('접두어가 올바르지 않습니다', 'INVALID_PREFIX')
    cases = import_workbook(upload, profiles, sheets, prefixes)
    for case in cases:
        case['source_refs'] = [ref.replace(upload.name, filename) for ref in case['source_refs']]
        case['import_origin'] = {'filename': 'xlsx:' + filename, 'sheet': case['sheet'], 'source_tc_id': case.get('source_tc_id', '')}
    record = {'preview_id': preview_id, 'filename': filename, 'sheets': sheets,
              'sha256': hashlib.sha256(upload.read_bytes()).hexdigest(),
              'meta_sha256': hashlib.sha256(upload.with_suffix('.json').read_bytes()).hexdigest(),
              'profiles': {sheet: profiles[sheet].to_dict() for sheet in sheets}}
    return record, cases, prefixes


def _identity(case: dict) -> tuple:
    ref = next((r for r in case.get('source_refs', []) if r.startswith('xlsx:')), '')
    origin = case.get('import_origin') or {}
    filename = origin.get('filename') or ref.split('#', 1)[0]
    sheet = origin.get('sheet') or (ref.split('#', 1)[1].rsplit('!R', 1)[0] if '#' in ref else case.get('sheet'))
    source_id = origin.get('source_tc_id', case.get('source_tc_id'))
    return (filename, sheet, source_id) if source_id else (ref,)


def _content(case: dict) -> dict:
    defaults = {'source_tc_id': '', 'tags': []}
    return {field: case.get(field, defaults.get(field)) for field in EDITABLE_FIELDS if field != 'status'}


def plan_import(body: dict) -> dict:
    suite = str(body.get('suite', '')).strip()
    suite_dir(suite)
    with suite_lock(suite):
        existing = load_cases(suite, include_deleted=True)
        by_identity = {}
        by_id = {case['case_id']: case for case in existing}
        for case in existing:
            by_identity.setdefault(_identity(case), []).append(case)
        records, incoming, prefixes = [], [], {}
        seen_sheets = set()
        for source in _sources(body):
            record, cases, source_prefixes = _read_source(source, body)
            if seen_sheets.intersection(record['sheets']):
                raise LibraryError('여러 파일의 동일 시트 이름은 함께 가져올 수 없습니다', 'DUPLICATE_SHEET', 409)
            seen_sheets.update(record['sheets'])
            records.append(record)
            incoming.extend(cases)
            prefixes.update(source_prefixes)
        identities = [_identity(case) for case in incoming]
        source_ids = [(case['sheet'], case.get('source_tc_id')) for case in incoming if case.get('source_tc_id')]
        used = set(by_id)
        rows = []
        for case in incoming:
            matches = by_identity.get(_identity(case), [])
            old = matches[0] if len(matches) == 1 else by_id.get(case['case_id'])
            # Generated IDs must not accidentally overwrite another workbook's case.
            if not matches and old and not any(ref in old.get('source_refs', []) for ref in case.get('source_refs', [])):
                old = None
            if old:
                case['case_id'] = old['case_id']
            elif case['case_id'] in used:
                case['case_id'] = next_case_id(prefixes.get(case['sheet']) or case['case_id'].rsplit('_', 1)[0], used)
            used.add(case['case_id'])
            status, reason = 'new', ''
            duplicate = identities.count(_identity(case)) > 1 or (case.get('source_tc_id') and source_ids.count((case['sheet'], case['source_tc_id'])) > 1)
            occupied_source_id = case.get('source_tc_id') and not matches and any(other.get('source_tc_id') == case['source_tc_id'] and not other.get('deleted') for other in existing)
            if duplicate or len(matches) > 1 or occupied_source_id:
                status, reason = 'error', '원본 ID 또는 출처가 중복됩니다'
            elif old:
                if _content(old) == _content(case) and not old.get('deleted'):
                    status = 'same'
                elif old.get('rev', 1) > 1 or old.get('deleted'):
                    status, reason = 'conflict', '기존 케이스가 수정되었습니다. 덮어쓰기 또는 제외를 선택하세요'
                else:
                    status = 'updated'
            rows.append({'case_id': case['case_id'], 'sheet': case['sheet'], 'source_tc_id': case.get('source_tc_id', ''),
                         'feature': case['feature'], 'status': status, 'reason': reason, 'before': old, 'after': case})
        run = {'run_id': 'libimp_' + secrets.token_hex(8), 'kind': 'xlsx_library', 'suite': suite,
               'created_at': datetime.now().isoformat(timespec='microseconds'), 'status': 'preview_ready', 'sources': records, 'rows': rows,
               'summary': {status: sum(row['status'] == status for row in rows) for status in ('new', 'updated', 'same', 'conflict', 'error')},
               'before_hash': _fingerprint(_snapshot(suite))}
        _save(run)
        return run


def _copy_style(cell, target, cache: dict) -> None:
    """다른 워크북의 셀 서식을 값으로 옮긴다. cell._style은 원본 워크북의 서식 표 번호라 그대로 쓰면
    대상 워크북에서 없는 번호(IndexError)나 엉뚱한 서식을 가리킨다. 같은 서식은 한 번만 변환한다."""
    from copy import copy as shallow_copy
    if not cell.has_style:
        return
    key = tuple(cell._style)
    if key not in cache:
        for name in ('font', 'border', 'fill', 'number_format', 'protection', 'alignment'):
            setattr(target, name, shallow_copy(getattr(cell, name)))
        cache[key] = shallow_copy(target._style)
    target._style = shallow_copy(cache[key])


def _merge_templates(run: dict, before: dict) -> tuple[bytes, dict]:
    import openpyxl
    from copy import copy as shallow_copy
    template = before['template.xlsx']
    first_source = run['sources'][0]
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(template))) if template else openpyxl.load_workbook(_paths.TC_LIBRARY_DIR / '_uploads' / (first_source['preview_id'] + '.xlsx'))
    profiles = json.loads(base64.b64decode(before['template_profile.json'])) if before['template_profile.json'] else {}
    try:
        for source in run['sources']:
            styles: dict = {}          # 워크북마다 서식 번호가 다르다 — 올린 파일별로 새로 변환한다
            uploaded = openpyxl.load_workbook(_paths.TC_LIBRARY_DIR / '_uploads' / (source['preview_id'] + '.xlsx'))
            try:
                if not template and source is first_source:
                    profiles.update(source['profiles'])
                    continue
                preserved_auxiliary = [name for name in uploaded.sheetnames if name not in wb.sheetnames and name not in source['sheets']]
                for sheet in source['sheets'] + preserved_auxiliary:
                    position = wb.sheetnames.index(sheet) if sheet in wb.sheetnames else len(wb.worksheets)
                    if sheet in wb.sheetnames:
                        wb.remove(wb[sheet])
                    src = uploaded[sheet]
                    dst = wb.create_sheet(sheet, position)
                    for row in src:
                        for cell in row:
                            target = dst.cell(cell.row, cell.column, cell.value)
                            _copy_style(cell, target, styles)
                            if cell.hyperlink:
                                target._hyperlink = shallow_copy(cell.hyperlink)
                            if cell.comment:
                                target.comment = shallow_copy(cell.comment)
                    for key, dimension in src.row_dimensions.items():
                        dst.row_dimensions[key] = shallow_copy(dimension)
                        dst.row_dimensions[key].parent = dst
                    for key, dimension in src.column_dimensions.items():
                        dst.column_dimensions[key] = shallow_copy(dimension)
                        dst.column_dimensions[key].parent = dst
                    for name in ('sheet_format', 'sheet_properties', 'page_margins', 'page_setup', 'print_options', 'data_validations', 'conditional_formatting', 'sheet_view', 'auto_filter'):
                        if name == 'sheet_view':
                            continue
                        setattr(dst, name, copy.deepcopy(getattr(src, name)))
                    dst.freeze_panes = src.freeze_panes
                    for merged in src.merged_cells.ranges:
                        dst.merge_cells(str(merged))
                    if sheet in source['profiles']:
                        profiles[sheet] = source['profiles'][sheet]
            finally:
                uploaded.close()
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue(), profiles
    finally:
        wb.close()


def commit_import(body: dict, actor: str = 'web') -> dict:
    run = get_run(body['run_id']) if body.get('run_id') else plan_import(body)
    with suite_lock(run['suite']):
        run = get_run(run['run_id'])
        if body.get('suite') and body['suite'] != run['suite']:
            raise LibraryError('다른 스위트의 작업입니다', 'INVALID_RUN')
        if run['status'] == 'committed':
            return {'run_id': run['run_id'], 'suite': run['suite'], 'status': 'committed', **run['result']}
        if run['status'] != 'preview_ready':
            raise LibraryError('반영할 수 없는 작업입니다', 'INVALID_RUN', 409)
        for source in run['sources']:
            path = _paths.TC_LIBRARY_DIR / '_uploads' / (source['preview_id'] + '.xlsx')
            if not path.exists() or hashlib.sha256(path.read_bytes()).hexdigest() != source['sha256'] or not path.with_suffix('.json').exists() or hashlib.sha256(path.with_suffix('.json').read_bytes()).hexdigest() != source['meta_sha256']:
                raise LibraryError('원본 파일이 변경되었습니다', 'SOURCE_CHANGED', 409)
        before = _snapshot(run['suite'])
        if _fingerprint(before) != run['before_hash']:
            raise LibraryError('스위트가 변경되었습니다. 다시 미리보세요', 'SUITE_CHANGED', 409)
        skip, overwrite = set(body.get('skip') or []), set(body.get('overwrite') or [])
        ids = {row['case_id'] for row in run['rows']}
        if (skip | overwrite) - ids or skip & overwrite:
            raise LibraryError('충돌 결정이 올바르지 않습니다', 'INVALID_DECISIONS')
        if any(row['status'] == 'error' and row['case_id'] not in skip for row in run['rows']):
            raise LibraryError('오류 행을 제외하세요', 'IMPORT_ERRORS', 409)
        if any(row['status'] == 'conflict' and row['case_id'] not in skip | overwrite for row in run['rows']):
            raise LibraryError('충돌마다 덮어쓰기 또는 제외를 선택하세요', 'UNRESOLVED_CONFLICT', 409)
        data = json.loads(base64.b64decode(before['cases.json'])) if before['cases.json'] else {'suite': run['suite'], 'sheets': [], 'cases': []}
        indexed = {case['case_id']: case for case in data['cases']}
        entries = []
        result = {'created': 0, 'updated': 0, 'unchanged': 0, 'skipped': len(skip)}
        for row in run['rows']:
            if row['case_id'] in skip:
                continue
            if row['status'] == 'same':
                result['unchanged'] += 1
                continue
            case = copy.deepcopy(row['after'])
            old = indexed.get(case['case_id'])
            if old:
                case = {**old, **_content(case), 'import_origin': case.get('import_origin', {}), 'deleted': False, 'rev': old['rev'] + 1, 'updated_at': now_iso()}
                data['cases'][data['cases'].index(old)] = case
                result['updated'] += 1
            else:
                data['cases'].append(case)
                result['created'] += 1
            entries.append(_entry(case['case_id'], '*', old, case, actor, 'import'))
        data['sheets'] = list(dict.fromkeys(data.get('sheets', []) + [sheet for source in run['sources'] for sheet in source['sheets']]))
        template, profiles = _merge_templates(run, before)
        after = {**before, 'cases.json': base64.b64encode(json.dumps(data, ensure_ascii=False, indent=2).encode()).decode(),
                 'template.xlsx': base64.b64encode(template).decode(),
                 'template_profile.json': base64.b64encode(json.dumps(profiles, ensure_ascii=False, indent=2).encode()).decode()}
        run.update(before_snapshot=before, result=result, decisions={'skip': sorted(skip), 'overwrite': sorted(overwrite)})
        _transaction(run, after, entries, 'committed')
        return {'run_id': run['run_id'], 'suite': run['suite'], 'status': 'committed', **result}


def rollback_import(run_id: str, actor: str = 'web') -> dict:
    run = get_run(run_id)
    with suite_lock(run['suite']):
        run = get_run(run_id)
        if run['status'] == 'rolled_back':
            return {'run_id': run_id, 'status': 'rolled_back'}
        if run['status'] != 'committed':
            raise LibraryError('반영된 작업만 복구할 수 있습니다', 'INVALID_RUN', 409)
        current = _snapshot(run['suite'])
        if _fingerprint(current) != run['after_hash']:
            raise LibraryError('반영 이후 수정되었습니다. 복구할 수 없습니다', 'SUITE_CHANGED', 409)
        after = copy.deepcopy(run['before_snapshot'])
        previous = json.loads(base64.b64decode(after['cases.json'])) if after['cases.json'] else {'suite': run['suite'], 'sheets': [], 'cases': []}
        current_cases = json.loads(base64.b64decode(current['cases.json']))['cases']
        old_by_id = {case['case_id']: case for case in previous['cases']}
        entries = []
        for case in current_cases:
            restored = old_by_id.get(case['case_id'])
            if restored is None:
                restored = {**case, 'deleted': True}
                previous['cases'].append(restored)
            restored['rev'] = max(restored.get('rev', 1), case['rev']) + 1
            restored['updated_at'] = now_iso()
            entries.append(_entry(case['case_id'], '*', case, copy.deepcopy(restored), actor, 'import_rollback'))
        after['cases.json'] = base64.b64encode(json.dumps(previous, ensure_ascii=False, indent=2).encode()).decode()
        _transaction(run, after, entries, 'rolled_back')
        return {'run_id': run_id, 'status': 'rolled_back'}
