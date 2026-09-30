"""엑셀 시트 → 라이브러리 케이스 (PRD F2.2, F2.4, F6.7)."""
from __future__ import annotations

import re
from pathlib import Path

from _tc_model import (
    EXCEL_TO_RESULT, merge_results, new_case, next_case_id, parse_steps, split_expected,
)
from _tc_template import TemplateProfile

_SYSTEM_LINE = re.compile(r"^id:(?P<id>[\w-]+)(?P<rest>(\s*\|\s*src:\S+)*)\s*$")


def split_note(text: str) -> tuple[str, str | None, list[str]]:
    """기타 칸 → (사람 메모, case_id, source_refs). 시스템 줄은 마지막 줄 `id:… | src:…`."""
    lines = (text or "").splitlines()
    if lines and (m := _SYSTEM_LINE.match(lines[-1].strip())):
        refs = re.findall(r"src:(\S+)", m.group("rest") or "")
        return "\n".join(lines[:-1]).strip(), m.group("id"), refs
    return (text or "").strip(), None, []


def _merged_lookup(ws) -> dict[tuple[int, int], object]:
    """병합 범위의 모든 칸 → 왼쪽 위 칸 값."""
    lookup = {}
    for rng in ws.merged_cells.ranges:
        top_left = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                lookup[(r, c)] = top_left
    return lookup


def read_sheet_cases(
    ws, profile: TemplateProfile, *, source_name: str, prefix: str,
    existing_ids: set[str],
) -> list[dict]:
    merged = _merged_lookup(ws)
    cols = profile.columns

    def value(row: int, key: str) -> str:
        col = cols.get(key)
        if col is None:
            return ""
        raw = merged.get((row, col), ws.cell(row, col).value)
        return "" if raw is None else str(raw).strip()

    used = set(existing_ids)
    cases: list[dict] = []
    path = ["", "", ""]
    feature = ""
    for r in range(profile.data_start_row, ws.max_row + 1):
        steps_text, expected_text = value(r, "steps"), value(r, "expected")
        raw_note = value(r, "note")
        if not (steps_text or expected_text or split_note(raw_note)[1]):
            continue  # 서식만 있는 빈 행. 라이브러리에서 온 행은 기타 칸의 id로 살린다
        for depth, key in enumerate(("l1", "l2", "l3")):
            cell = value(r, key)
            if cell and cell != path[depth]:
                path[depth] = cell
                for deeper in range(depth + 1, 3):
                    path[deeper] = ""
                feature = ""
        feature = value(r, "feature") or feature
        note, case_id, refs = split_note(raw_note)
        if not case_id or case_id in used:
            case_id = next_case_id(prefix, used)
        used.add(case_id)
        expected, bullets = split_expected(expected_text)
        results = [
            EXCEL_TO_RESULT.get(str(ws.cell(r, c).value or "").strip().lower(), "")
            for c in profile.result_columns.values()
        ]
        cases.append(new_case(
            case_id=case_id, sheet=profile.sheet, path=list(path), feature=feature,
            precondition=value(r, "precondition"), steps=parse_steps(steps_text),
            expected=expected, bullets=bullets, priority=value(r, "priority"),
            auto=value(r, "auto"), execution_result=merge_results(results),
            status="approved", note=note,
            source_refs=refs or [f"xlsx:{source_name}#{profile.sheet}!R{r}"],
        ))
    return cases


def import_workbook(
    path: Path, profiles: dict[str, TemplateProfile], sheets: list[str],
    prefixes: dict[str, str],
) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=False)
    try:
        cases: list[dict] = []
        used: set[str] = set()
        for index, sheet in enumerate(sheets, 1):
            prefix = prefixes.get(sheet) or f"S{index:02d}"
            sheet_cases = read_sheet_cases(
                wb[sheet], profiles[sheet], source_name=path.name, prefix=prefix,
                existing_ids=used,
            )
            used.update(c["case_id"] for c in sheet_cases)
            cases.extend(sheet_cases)
        return cases
    finally:
        wb.close()
