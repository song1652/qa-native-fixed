"""dash_excel.py — Excel Import 유틸리티 (serve.py Phase-1 분리).

serve.py의 DashboardHandler가 직접 import해서 사용하는 순수 함수 모음.
외부 상태에 의존하지 않으므로 단독으로 테스트 가능하다.
"""
from __future__ import annotations

import re
from pathlib import Path

from _paths import IMPORT_DIR


# ── Excel Import 유틸 ─────────────────────────────────────────

def _list_import_files() -> list:
    """import/ 폴더의 .xlsx 파일 목록."""
    if not IMPORT_DIR.exists():
        return []
    return sorted([f.name for f in IMPORT_DIR.glob("*.xlsx")])


def _detect_header_row(ws) -> int | None:
    """'Test\\nScenario ID' 패턴이 있는 헤더 행 번호 반환."""
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=False), 1):
        for cell in row:
            if cell.value and "Scenario ID" in str(cell.value):
                return i
    return None


def _list_excel_sheets(filepath: Path) -> list:
    """엑셀 파일의 시트별 테스트케이스 수 반환."""
    import openpyxl
    wb = openpyxl.load_workbook(str(filepath), data_only=True, read_only=True)
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        header_row = _detect_header_row(ws)
        if header_row is None:
            continue
        count = 0
        for row in ws.iter_rows(min_row=header_row + 2, values_only=False):
            cell_b = row[1].value if len(row) > 1 else None
            if cell_b and "_" in str(cell_b):
                count += 1
        if count > 0:
            sheets.append({"name": name, "count": count})
    wb.close()
    return sheets


def _parse_excel_sheet(wb, sheet_name: str) -> list:
    """엑셀 시트에서 테스트케이스 목록 추출."""
    ws = wb[sheet_name]
    header_row = _detect_header_row(ws)
    if header_row is None:
        return []

    last_main = ""
    last_sub = ""
    cases = []
    for row in ws.iter_rows(min_row=header_row + 2, values_only=False):
        tc_id = row[1].value if len(row) > 1 else None
        if not tc_id or "_" not in str(tc_id):
            continue
        main = str(row[2].value).strip() if len(row) > 2 and row[2].value else ""
        sub = str(row[3].value).strip() if len(row) > 3 and row[3].value else ""
        detail = str(row[4].value).strip() if len(row) > 4 and row[4].value else ""
        summary = str(row[5].value).strip() if len(row) > 5 and row[5].value else ""
        precond = str(row[6].value).strip() if len(row) > 6 and row[6].value else ""
        steps = str(row[7].value).strip() if len(row) > 7 and row[7].value else ""
        expected = str(row[8].value).strip() if len(row) > 8 and row[8].value else ""
        level = str(row[9].value).strip() if len(row) > 9 and row[9].value else ""
        if main:
            last_main = main
        else:
            main = last_main
        if sub:
            last_sub = sub
        else:
            sub = last_sub
        cases.append({
            "main": main, "sub": sub, "detail": detail,
            "summary": summary, "precondition": precond,
            "steps": steps, "expected": expected, "level": level,
        })
    return cases


def _level_to_priority(level: str) -> str:
    level = level.strip()
    if level in ("BAT", "Level 1"):
        return "high"
    elif level == "Level 2":
        return "medium"
    return "low"


def _to_slug(text: str) -> str:
    """TC Summary → 파일명 슬러그."""
    if not text:
        return "unnamed"
    text = text.replace("\n", " ").strip()
    text = re.sub(r'[/\\:*?"<>|.\[\]()>{},]', '', text)
    text = re.sub(r'\s+', '_', text)
    text = re.sub(r'_+', '_', text).strip('_')
    return text[:60]


def _write_tc_files(cases: list, output_dir: Path) -> int:
    """케이스 목록을 tc_*.md 파일로 생성. 생성 건수 반환."""
    output_dir.mkdir(parents=True, exist_ok=True)
    for i, c in enumerate(cases):
        num = str(i + 1).zfill(3)
        slug = _to_slug(c["summary"])
        filepath = output_dir / f"tc_{num}_{slug}.md"
        priority = _level_to_priority(c.get("level", ""))
        tags = []
        if c["main"]:
            clean = re.sub(r'\(.*?\)', '', c["main"]).strip().replace('\n', '')
            if clean:
                tags.append(clean)
        if c["sub"]:
            tags.append(c["sub"].replace('\n', ''))
        if not tags:
            tags = ["general"]
        tags_str = ", ".join(tags)
        steps_lines = [s.strip() for s in c["steps"].split("\n")
                       if s.strip() and not s.strip().startswith("0.")]
        steps_text = "\n".join(steps_lines) if steps_lines else "1. (스텝 미기재)"
        exp_lines = []
        for e in c["expected"].split("\n"):
            e = e.strip()
            if e:
                if not e.startswith("-") and not e.startswith("*"):
                    e = f"- {e}"
                exp_lines.append(e)
        expected_text = "\n".join(exp_lines) if exp_lines else "- (기대결과 미기재)"
        pre_lines = [p.strip() for p in c["precondition"].split("\n") if p.strip()]
        precond_text = "\n".join(pre_lines) if pre_lines else "- 없음"
        title = c["summary"].replace("\n", " ").strip()
        content = (
            f"---\nid: tc_{num}\ndata_key: null\npriority: {priority}\n"
            f"tags: [{tags_str}]\ntype: structured\n---\n"
            f"# {title}\n\n## Precondition\n{precond_text}\n\n"
            f"## Steps\n{steps_text}\n\n## Expected\n{expected_text}\n"
        )
        filepath.write_text(content, encoding="utf-8")
    return len(cases)
