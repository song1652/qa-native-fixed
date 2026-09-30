"""TC 스튜디오 실제 브라우저 검증 (로드맵 W1~W7). 실행: .venv/bin/python -m pytest tests/unit/tc_library/test_tc_studio_e2e.py"""
from __future__ import annotations

import json
import re
import urllib.request
from pathlib import Path
from urllib.parse import quote

import pytest
from playwright.sync_api import Page, expect

from tests.unit.import_studio.import_studio_test_support import dashboard_server, request_json
from tests.unit.tc_library.tc_fixtures import build_template_workbook

S = quote("야핏무브")


def _seed(base_url: str, tmp_path: Path, *, suite="야핏무브", sheets=None) -> None:
    data = build_template_workbook(tmp_path / "seed.xlsx").read_bytes()
    req = urllib.request.Request(
        base_url + "/api/tc-library/import/preview?filename=" + quote("야핏무브_Full.xlsx"),
        data=data, method="POST", headers={"Content-Type": "application/octet-stream"})
    with urllib.request.urlopen(req, timeout=10) as resp:
        preview = json.loads(resp.read())
    status, body = request_json(base_url, "POST", "/api/tc-library/import", {
        "preview_id": preview["preview_id"], "suite": suite, "sheets": sheets or ["혜택", "홈"],
        "prefixes": {"혜택": "BEN", "홈": "HOME"}})
    assert status == 200, body


@pytest.fixture
def studio(tmp_path: Path, page: Page):
    import _paths
    original_paths = {name: getattr(_paths, name) for name in ("PROJECT_ROOT", "TESTCASES_DIR", "TC_LIBRARY_DIR")}
    project = tmp_path / "project"
    (project / "testcases").mkdir(parents=True)
    with dashboard_server(project) as base_url:
        _seed(base_url, tmp_path)
        page.goto(base_url + "/tc-studio")
        expect(page.locator("#grid-body tr[data-case]")).to_have_count(6)
        yield base_url, page
    assert {name: getattr(_paths, name) for name in original_paths} == original_paths


def rows(page: Page):
    return page.locator("#grid-body tr[data-case]")


# ── W1: 셸·라이브러리·상세 ──────────────────────────────────────────────────────────
def test_route_sidebar_and_tabs(studio):
    _, page = studio
    expect(page.locator("#tab-tc_studio")).to_have_class(re.compile("active"))
    expect(page.locator('[data-id="suite-select"] option')).to_have_text(["야핏무브 (6)"])
    expect(page.locator('[data-id="nav-tab-library"]')).to_be_visible()
    expect(page.locator('[data-id="nav-tab-generate"]')).to_be_visible()
    expect(page.locator('[data-id="nav-tab-review"]')).to_be_visible()


def test_existing_suite_selection_clears_previous_branch_and_loads_cases(studio, tmp_path):
    base, page = studio
    _seed(base, tmp_path, suite="다른스위트", sheets=["홈"])
    page.reload()
    page.locator('[data-id="tree-node"][data-name="상단 배너"]').click()
    expect(rows(page)).to_have_count(2)
    page.locator('[data-id="suite-select"]').select_option("다른스위트")
    expect(rows(page)).to_have_count(1)
    expect(page.locator('[data-id="import-modal"]')).to_be_hidden()
    expect(page.locator('[data-id="nav-tab-library"]')).to_have_attribute("aria-selected", "true")
    page.locator('[data-id="nav-tab-generate"]').click()
    expect(page.locator('[data-id="gen-target-sheet"] option')).to_have_text(["시트를 선택하세요", "홈"])
    expect(page.locator('#job')).to_be_hidden()
    expect(page.locator('#tcs-toasts')).not_to_contain_text('요청을 처리하지 못했습니다')


def test_tree_filters_and_search(studio):
    _, page = studio
    page.locator('[data-id="tree-node"][data-name="상단 배너"]').click()
    expect(rows(page)).to_have_count(2)
    expect(page.locator("#grid-crumb")).to_have_text("혜택 › 혜택 탭 › 상단 배너")
    page.locator('[data-id="lib-filter-reset"]').click()
    page.locator('[data-id="lib-search"]').fill("돈불리기")
    expect(rows(page)).to_have_count(2)
    page.locator('[data-id="lib-filter-reset"]').click()
    page.locator('[data-id="lib-filter-result"]').select_option("fail")
    expect(rows(page)).to_have_count(1)
    page.locator('[data-id="lib-filter-result"]').select_option("none")
    expect(rows(page)).to_have_count(3)


def test_chip_and_cell_edits_persist(studio):
    _, page = studio
    row = page.locator('tr[data-case="BEN_0002"]')
    headers = page.locator("#grid thead th").all_text_contents()
    assert headers[6] == "E제목"
    priority_index = next(i for i, text in enumerate(headers) if "우선순위" in text)
    assert headers[priority_index + 1] == "실행 결과"
    expect(row.locator("td").nth(priority_index + 1).locator('[data-id="grid-result"]')).to_be_visible()
    row.locator('[data-id="grid-cell-priority"]').select_option("P3")
    expect(page.locator(".toast.ok").first).to_contain_text("BEN_0002 저장됨")
    cell = row.locator('[data-id="grid-cell-feature"]')
    cell.dblclick()
    page.keyboard.press("ControlOrMeta+a")
    page.keyboard.type("배너 가로 스크롤")
    page.keyboard.press("ControlOrMeta+Enter")
    expect(page.locator(".toast.ok").last).to_contain_text("rev 3")
    page.reload()
    row = page.locator('tr[data-case="BEN_0002"]')
    expect(row.locator('[data-id="grid-cell-priority"]')).to_have_value("P3")
    expect(row.locator('[data-id="grid-cell-feature"]')).to_have_text("배너 가로 스크롤")


def test_stale_rev_shows_conflict_toast(studio):
    base_url, page = studio
    request_json(base_url, "PATCH", f"/api/tc-library/{S}/cases/BEN_0003", {"rev": 1, "priority": "P0"})
    page.locator('tr[data-case="BEN_0003"] [data-id="grid-cell-priority"]').select_option("P2")
    expect(page.locator('[data-id="toast-conflict-reload"]')).to_be_visible()
    page.locator('[data-id="toast-conflict-reload"]').click()
    expect(page.locator('tr[data-case="BEN_0003"] [data-id="grid-cell-priority"]')).to_have_value("P0")


def test_bulk_add_delete_and_undo(studio):
    _, page = studio
    rows(page).nth(0).locator('[data-id="grid-row-check"]').check()
    rows(page).nth(1).locator('[data-id="grid-row-check"]').check()
    expect(page.locator('[data-id="bulk-bar"]')).to_contain_text("2건 선택")
    page.locator('[data-id="bulk-priority"]').select_option("P2")
    expect(page.locator('tr[data-case="BEN_0002"] [data-id="grid-cell-priority"]')).to_have_value("P2")

    page.locator('[data-id="btn-add-case"]').click()
    expect(rows(page)).to_have_count(7)
    expect(page.locator('[data-id="detail-panel"] #d-id')).to_have_text("BEN_0006")

    page.locator('[data-id="detail-delete"]').click()
    page.locator('[data-id="confirm-ok"]').click()
    expect(rows(page)).to_have_count(6)
    page.locator('[data-id="delete-undo"]').click()
    expect(rows(page)).to_have_count(7)


def test_detail_save_history_and_revert(studio):
    _, page = studio
    # Phase 3의 배너 응답을 늦춰 저장 중 선택한 이력 탭이 유지되는지 재현한다.
    page.evaluate("""() => {
        const watch = window.TCS_NS.sourceWatch;
        if (!watch) return;
        const original = watch.renderBanner;
        watch.renderBanner = async (...args) => {
            await new Promise((resolve) => setTimeout(resolve, 1000));
            return original(...args);
        };
    }""")
    page.locator('tr[data-case="HOME_0001"] td.no').click()
    expect(page.locator("#d-id")).to_have_text("HOME_0001")
    page.locator('[data-id="detail-expected"]').fill("홈 화면 상단에 D+3 전용 카드가 노출된다.")
    page.locator('[data-id="detail-step-add"]').click()
    page.locator('[data-id="detail-step-input"]').last.fill("홈 카드 선택")
    page.locator('[data-id="detail-save"]').click()
    expect(page.locator("#d-rev")).to_have_text("rev 2")
    page.locator('[data-id="detail-tab-history"]').click()
    history = page.locator('[data-id="detail-history"] li')
    expect(history.first).to_contain_text("expected")
    expect(page.locator('[data-id="detail-save"]')).not_to_have_class(re.compile("loading"))
    page.locator('[data-id="detail-history-revert"]').first.click()
    expect(page.locator("#d-rev")).to_have_text("rev 3")


# ── W2: 엑셀 가져오기 모달 ───────────────────────────────────────────────
def test_import_modal_previews_and_imports_workbook(tmp_path, page):
    project = tmp_path / "project"
    (project / "testcases").mkdir(parents=True)
    with dashboard_server(project) as base_url:
        page.goto(base_url + "/tc-studio")
        page.locator('[data-id="empty-import-xlsx"]').click()
        expect(page.locator('[data-id="import-modal"]')).to_be_visible()
        assert page.evaluate("""async () => {
            const modal = document.querySelector('[data-id="import-modal"]');
            await refreshAll();
            return modal === document.querySelector('[data-id="import-modal"]') && !modal.hidden;
        }""")
        xlsx = build_template_workbook(tmp_path / "야핏무브_Full.xlsx")
        page.locator('[data-id="import-file-input"]').set_input_files(xlsx)
        expect(page.locator('[data-id="import-confirm"]')).to_have_text("6건 가져오기")
        expect(page.locator('[data-id="import-sheets"] input[type=checkbox]')).to_have_count(2)
        page.locator('[data-prefix="혜택"]').fill("BEN")
        page.locator('[data-prefix="홈"]').fill("HOME")
        page.locator('[data-id="import-plan"]').click()
        page.locator('[data-id="import-confirm"]').click()
        expect(page.locator('[data-id="import-modal"]')).to_be_hidden()
        expect(rows(page)).to_have_count(6)
        expect(page.locator('[data-id="suite-select"] option')).to_have_text(["야핏무브 (6)"])
        expect(page.locator('tr[data-case="BEN_0001"]')).to_be_visible()
        page.reload()
        expect(rows(page)).to_have_count(6)

# ── W3: 내보내기 화면 ──────────────────────────────────────────────────────────
def test_export_check_and_download(studio):
    _, page = studio
    page.locator('[data-id="nav-tab-export"]').click()
    page.locator('[data-id="xlsx-run-check"]').click()
    expect(page.locator('[data-id="xlsx-integrity"] li .ok')).to_have_count(4)
    expect(page.locator('[data-id="xlsx-download"]')).to_be_enabled()
    with page.expect_download() as info:
        page.locator('[data-id="xlsx-download"]').click()
    assert info.value.suggested_filename.endswith(".xlsx")


def test_blank_template_can_start_generation(studio, tmp_path):
    """A template without TC content must still expose its sheets for generation."""
    import openpyxl
    base_url, page = studio
    template = build_template_workbook(tmp_path / 'blank.xlsx')
    wb = openpyxl.load_workbook(template)
    for ws in [wb['혜택'], wb['홈']]:
        for merged in list(ws.merged_cells.ranges):
            if merged.min_row >= 13:
                ws.unmerge_cells(str(merged))
        for cells in ws.iter_rows(min_row=13):
            for cell in cells:
                cell.value = None
    wb.save(template)
    wb.close()
    page.locator('[data-id="btn-import-xlsx"]').click()
    page.locator('[data-id="import-file-input"]').set_input_files(template)
    expect(page.locator('[data-id="import-confirm"]')).to_have_text('0건 가져오기')
    page.locator('[data-id="import-suite"]').fill('빈양식')
    page.locator('[data-id="import-plan"]').click()
    page.locator('[data-id="import-confirm"]').click()
    expect(page.locator('[data-id="suite-select"]')).to_have_value('빈양식')
    expect(rows(page)).to_have_count(0)
    page.locator('[data-id="nav-tab-generate"]').click()
    expect(page.locator('[data-id="gen-target-sheet"]')).to_have_value('')
    expect(page.locator('[data-id="gen-target-sheet"] option')).to_have_text(['시트를 선택하세요', '혜택', '홈'])
    page.locator('[data-id="gen-target-sheet"]').select_option('혜택')
    expect(page.locator('[data-id="gen-new-l1"]')).to_be_visible()
    page.locator('[data-id="gen-new-l1"]').fill('사용자 기능')
    page.locator('[data-id="gen-path-l2"]').select_option('__new')
    page.locator('[data-id="gen-new-l2"]').fill('로그인')
    page.locator('[data-id="gen-path-l3"]').select_option('__new')
    page.locator('[data-id="gen-new-l3"]').fill('오류 처리')
    page.locator('[data-id="gen-add-branch"]').click()
    expect(page.locator('[data-id="gen-path-l1"]')).to_have_value('사용자 기능')
    expect(page.locator('[data-id="gen-path-l2"]')).to_have_value('로그인')
    expect(page.locator('[data-id="gen-path-l3"]')).to_have_value('오류 처리')
    page.reload()
    page.locator('[data-id="gen-target-sheet"]').select_option('혜택')
    expect(page.locator('[data-id="gen-path-l1"] option')).to_contain_text(['사용자 기능', '+ 새로 만들기…'])
    page.locator('[data-id="gen-path-l2"]').select_option('로그인')
    page.locator('[data-id="gen-path-l3"]').select_option('오류 처리')
    expect(page.locator('[data-id="gen-style-examples"]')).to_contain_text('입력한 정보와 작성 규칙')
    page.locator('[data-id="src-tab-paste"]').click()
    page.locator('[data-id="src-paste"]').fill('초대 링크 복사 버튼을 누르면 링크가 클립보드에 복사된다.')
    page.locator('[data-id="src-paste-add"]').click()
    expect(page.locator('[data-id="gen-submit"]')).to_be_enabled()
    _, data = request_json(base_url, 'GET', '/api/tc-library/' + quote('빈양식'))
    assert data['total'] == 0
    page.reload()
    expect(page.locator('[data-id="nav-tab-generate"]')).to_have_attribute('aria-selected', 'true')
    expect(page.locator('.wizard [role="tab"]').first).to_have_attribute('data-id', 'nav-tab-generate')
    expect(page.locator('[data-id="src-file-drop"]')).to_be_visible()
    page.locator('[data-id="src-tab-url"]').click()
    expect(page.locator('[data-id="src-url-url"]')).to_be_visible()


def test_sheet_name_can_be_changed_and_survives_reload(studio):
    base_url, page = studio
    page.locator('[data-id="nav-tab-generate"]').click()
    expect(page.locator('[data-id="gen-target-sheet"]')).to_have_value('')
    page.locator('[data-id="gen-target-sheet"]').select_option('혜택')
    native_dialogs = []
    page.on('dialog', lambda dialog: (native_dialogs.append(dialog.message), dialog.dismiss()))
    page.locator('[data-id="gen-rename-sheet"]').click()
    expect(page.locator('[data-id="sheet-rename-modal"]')).to_be_visible()
    expect(page.locator('[data-id="sheet-rename-name"]')).to_be_focused()
    expect(page.locator('[data-id="sheet-rename-name"]')).to_have_value('혜택')
    page.locator('[data-id="sheet-rename-name"]').fill('회원 혜택')
    page.locator('[data-id="sheet-rename-save"]').click()
    expect(page.locator('[data-id="gen-target-sheet"]')).to_have_value('회원 혜택')
    assert native_dialogs == []
    expect(page.locator('[data-id="sheet-rename-modal"]')).not_to_be_visible()
    page.reload()
    page.locator('[data-id="nav-tab-generate"]').click()
    expect(page.locator('[data-id="gen-target-sheet"] option')).to_have_text(['시트를 선택하세요', '회원 혜택', '홈'])
    _, body = request_json(base_url, 'GET', f'/api/tc-library/{S}/cases/BEN_0001')
    assert body['case']['sheet'] == '회원 혜택'


def test_new_sheet_and_empty_branch_can_be_added(studio):
    base_url, page = studio
    page.locator('[data-id="nav-tab-generate"]').click()
    page.locator('[data-id="gen-add-sheet"]').click()
    expect(page.locator('[data-id="sheet-rename-modal"]')).to_be_visible()
    page.locator('[data-id="sheet-rename-name"]').press('Escape')
    expect(page.locator('[data-id="sheet-rename-modal"]')).not_to_be_visible()
    expect(page.locator('[data-id="gen-add-sheet"]')).to_be_focused()
    page.locator('[data-id="gen-add-sheet"]').click()
    page.locator('[data-id="sheet-rename-name"]').fill('혜택')
    page.locator('[data-id="sheet-rename-save"]').click()
    expect(page.locator('[data-id="sheet-rename-error"]')).to_contain_text('이미 존재')
    expect(page.locator('[data-id="sheet-rename-modal"]')).to_be_visible()
    page.locator('[data-id="sheet-rename-name"]').fill('회원가입')
    page.locator('[data-id="sheet-rename-save"]').click()
    expect(page.locator('[data-id="gen-target-sheet"]')).to_have_value('회원가입')
    expect(page.locator('[data-id="gen-new-l1"]')).to_have_value('')
    page.locator('[data-id="gen-new-l1"]').fill('이메일 가입')
    page.locator('[data-id="gen-add-branch"]').click()
    expect(page.locator('[data-id="gen-path-l1"]')).to_have_value('이메일 가입')
    page.reload()
    page.locator('[data-id="nav-tab-generate"]').click()
    page.locator('[data-id="gen-target-sheet"]').select_option('회원가입')
    expect(page.locator('[data-id="gen-path-l1"]')).to_have_value('이메일 가입')
    _, body = request_json(base_url, 'GET', f'/api/tc-library/{S}')
    assert body['total'] == 6 and not any(c['sheet'] == '회원가입' for c in body['items'])


def test_reentering_studio_with_empty_suite_opens_planning_screen(studio):
    """다른 메뉴에 갔다 돌아와도 0건 스위트면 라이브러리가 아니라 기획 정보 화면이 열린다."""
    from _tc_library import import_cases
    base_url, page = studio
    import_cases("빈스위트", ["테스트케이스"], [], "tester")
    page.reload()
    expect(page.locator("#grid-body tr[data-case]")).to_have_count(6)       # 첫 로딩이 끝난 뒤 스위트를 바꾼다
    page.locator("#suite-select").select_option("빈스위트")
    expect(page.locator('[data-screen="generate"].screen')).to_have_class(re.compile("active"))
    page.evaluate("selectView('reports')")
    page.evaluate("""() => { window.__screens = [];
      new MutationObserver(() => window.__screens.push(...[...document.querySelectorAll('.screen.active')].map(s => s.dataset.screen)))
        .observe(document.body, {subtree: true, childList: true, attributes: true, attributeFilter: ['class']}); }""")
    page.evaluate("selectView('tc_studio')")
    expect(page.locator("#suite-select")).to_have_value("빈스위트")
    expect(page.locator('[data-screen="generate"].screen')).to_have_class(re.compile("active"))
    assert "library" not in page.evaluate("window.__screens")      # 라이브러리가 잠깐 보였다 바뀌지 않는다
    expect(page.locator('[data-id="studio-loading"]')).to_be_hidden()    # 불러오기가 끝나면 안내는 사라진다


def test_export_selected_sheets_are_checkbox_chips_with_counts(studio):
    _, page = studio
    page.locator('[data-id="nav-tab-export"]').click()
    picks = page.locator('[data-id="xlsx-sheets"]')
    expect(picks).to_be_hidden()                                  # '선택한 시트'일 때만 펼친다
    page.locator('input[name="xscope"][value="sheets"]').check()
    expect(picks.locator("label")).to_have_count(2)
    expect(picks).to_contain_text("혜택")
    expect(page.locator("#x-sheets-n")).to_have_text("2/2개")
    counts = [int(t) for t in picks.locator(".n").all_inner_texts()]
    assert sum(counts) == 6
    picks.locator("label", has_text="홈").click()
    expect(page.locator("#x-sheets-n")).to_have_text("1/2개")
    picks.locator("label", has_text="혜택").click()
    expect(page.locator('[data-id="xlsx-run-check"]')).to_be_disabled()   # 하나도 안 고르면 검사 불가
    picks.locator("label", has_text="혜택").click()
    page.locator('[data-id="xlsx-run-check"]').click()
    expect(page.locator('[data-id="xlsx-download"]')).to_be_enabled()


def test_imported_cases_show_review_as_disabled(studio):
    _, page = studio
    rows(page).first.click()
    expect(page.locator('[data-id="detail-imported"]')).to_have_text("가져옴")
    expect(page.locator('[data-id="detail-status"]')).to_be_disabled()
    expect(page.locator('[data-id="detail-status"]')).to_have_attribute("title", re.compile("엑셀에서 가져온"))
    page.locator('[data-id="nav-tab-export"]').click()
    expect(page.locator('[data-id="xlsx-approved-note"]')).to_have_text("가져온 케이스 6건 제외")
