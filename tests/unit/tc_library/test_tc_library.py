from __future__ import annotations

import pytest

import _tc_library as lib
from _tc_template import analyze_workbook
from _tc_xlsx_import import import_workbook

SUITE = "야핏무브"


@pytest.fixture
def seeded(library_dir, template_xlsx):
    profiles = analyze_workbook(template_xlsx)
    cases = import_workbook(template_xlsx, profiles, ["혜택", "홈"], {"혜택": "BEN", "홈": "HOME"})
    lib.save_template(SUITE, template_xlsx, profiles)
    lib.import_cases(SUITE, ["혜택", "홈"], cases, "tester")
    return cases


def test_import_is_idempotent_and_updates_by_case_id(seeded):
    assert lib.import_cases(SUITE, ["혜택", "홈"], seeded, "tester") == \
        {"created": 0, "updated": 0, "unchanged": 6}
    changed = [dict(c) for c in seeded]
    changed[0]["priority"] = "P1"
    assert lib.import_cases(SUITE, ["혜택", "홈"], changed, "tester")["updated"] == 1
    assert lib.get_case(SUITE, "BEN_0001")["rev"] == 2
    assert lib.list_suites() == [{"suite": SUITE, "sheets": ["혜택", "홈"], "count": 6}]


def test_patch_bumps_rev_and_rejects_stale_rev(seeded):
    case = lib.patch_case(SUITE, "BEN_0002", 1, {"priority": "P1"}, "tester")
    assert (case["rev"], case["priority"]) == (2, "P1")

    with pytest.raises(lib.RevConflict) as exc:
        lib.patch_case(SUITE, "BEN_0002", 1, {"priority": "P2"}, "tester")
    assert exc.value.server_case["rev"] == 2 and exc.value.status == 409

    with pytest.raises(lib.LibraryError) as exc:
        lib.patch_case(SUITE, "BEN_0002", 2, {"rev": 99}, "tester")
    assert exc.value.code == "FIELD_NOT_EDITABLE"


def test_bulk_patch_applies_matching_revs_and_reports_conflicts(seeded):
    result = lib.bulk_patch(SUITE, [{"case_id": "BEN_0001", "rev": 1},
                                    {"case_id": "BEN_0002", "rev": 7}], {"auto": "Y-app"}, "t")
    assert result["updated"] == ["BEN_0001"]
    assert [c["case_id"] for c in result["conflicts"]] == ["BEN_0002"]


def test_history_and_revert(seeded):
    lib.patch_case(SUITE, "BEN_0002", 1, {"priority": "P1"}, "tester")
    entry = next(e for e in lib.history(SUITE, "BEN_0002") if e["field"] == "priority")
    assert (entry["before"], entry["after"], entry["actor"]) == ("", "P1", "tester")

    reverted = lib.revert(SUITE, "BEN_0002", entry["history_id"], 2, "tester")
    assert (reverted["priority"], reverted["rev"]) == ("", 3)


def test_create_duplicate_delete_restore(seeded):
    created = lib.create_case(SUITE, {"sheet": "혜택", "feature": "새 기능",
                                      "path": ["혜택 탭", "상단 배너", ""]}, "t", after="BEN_0003")
    assert (created["case_id"], created["status"]) == ("BEN_0006", "draft")
    order = [c["case_id"] for c in lib.load_cases(SUITE)]
    assert order.index("BEN_0006") == order.index("BEN_0003") + 1

    dup = lib.duplicate_case(SUITE, "BEN_0001", "t")
    assert (dup["case_id"], dup["feature"], dup["status"]) == ("BEN_0007", "혜택 탭 버튼", "draft")
    assert dup["source_refs"] == lib.get_case(SUITE, "BEN_0001")["source_refs"]

    assert lib.delete_cases(SUITE, [{"case_id": "BEN_0001", "rev": 1}], "t")["deleted"] == ["BEN_0001"]
    assert "BEN_0001" not in [c["case_id"] for c in lib.load_cases(SUITE)]
    assert lib.restore_case(SUITE, "BEN_0001", "t")["rev"] == 3


def test_tree_counts_and_filters(seeded):
    lib.create_case(SUITE, {"sheet": "혜택", "feature": "새 기능",
                            "path": ["혜택 탭", "상단 배너", ""]}, "t")
    tree = lib.build_tree(lib.load_cases(SUITE))
    benefit = tree[0]
    assert (benefit["name"], benefit["level"], benefit["count"], benefit["draft"]) == ("혜택", "sheet", 6, 1)
    banner = benefit["children"][0]["children"][1]
    assert (banner["name"], banner["level"], banner["count"], banner["invalid"]) == ("상단 배너", "l2", 3, 1)

    cases = lib.load_cases(SUITE)
    ids = lambda q: [c["case_id"] for c in lib.filter_cases(cases, q)]
    assert ids({"path": "혜택/혜택 탭/상단 배너"}) == ["BEN_0002", "BEN_0003", "BEN_0006"]
    assert ids({"q": "돈불리기"}) == ["BEN_0004", "BEN_0005"]
    assert ids({"invalid": "1"}) == ["BEN_0006"]
    assert ids({"execution_result": "fail"}) == ["BEN_0003"]
    assert ids({"priority": "-", "sheet": "홈"}) == []


def test_suite_name_cannot_escape_library_dir(library_dir):
    for bad in ("../etc", "_uploads", ""):
        with pytest.raises(lib.LibraryError):
            lib.suite_dir(bad)


def test_rename_sheet_preserves_cases_template_and_md_mapping(seeded):
    import openpyxl
    from _state import read_state, update_state
    root = lib.suite_dir(SUITE)
    update_state(root / 'md_export.json', lambda _: {'groups': [{'path': ['혜택', '혜택 탭'], 'group': 'benefits', 'code': 'BEN'}], 'ids': {'BEN_0001': 'BEN_01'}})
    jobs = root.parent / '_jobs' / 'job_rename' / 'status.json'
    update_state(jobs, lambda _: {'suite': SUITE, 'status': 'done', 'target': {'sheet': '혜택'}})
    lib.rename_sheet(SUITE, '혜택', '회원 혜택', 'tester')
    case = lib.get_case(SUITE, 'BEN_0001')
    assert case['sheet'] == '회원 혜택' and case['rev'] == 2
    assert case['feature'] == seeded[0]['feature']
    assert lib.list_suites()[0]['sheets'] == ['회원 혜택', '홈']
    assert lib.load_profiles(SUITE)['회원 혜택'].sheet == '회원 혜택'
    wb = openpyxl.load_workbook(root / 'template.xlsx')
    assert '회원 혜택' in wb.sheetnames and '혜택' not in wb.sheetnames
    wb.close()
    assert read_state(jobs)['target']['sheet'] == '회원 혜택'
    cfg = read_state(root / 'md_export.json')
    assert cfg['groups'][0]['path'] == ['회원 혜택', '혜택 탭']
    assert cfg['ids']['BEN_0001'] == 'BEN_01'
    assert lib.history(SUITE, 'BEN_0001')[0]['field'] == 'sheet'
    for invalid in ['홈', 'History', '', 'bad/name', 'x' * 32]:
        with pytest.raises(lib.LibraryError):
            lib.rename_sheet(SUITE, '회원 혜택', invalid, 'tester')
    assert lib.get_case(SUITE, 'BEN_0001')['sheet'] == '회원 혜택'


def test_rename_sheet_restores_template_when_case_write_fails(seeded, monkeypatch):
    from _state import read_state
    root = lib.suite_dir(SUITE)
    before_xlsx = (root / 'template.xlsx').read_bytes()
    before_profiles = read_state(root / 'template_profile.json')
    real_update = lib.update_state
    failed = False
    def fail_once(path, mutate):
        nonlocal failed
        if path.name == 'cases.json' and not failed:
            failed = True
            raise OSError('disk write failed')
        return real_update(path, mutate)
    monkeypatch.setattr(lib, 'update_state', fail_once)
    with pytest.raises(OSError, match='disk write failed'):
        lib.rename_sheet(SUITE, '혜택', '회원 혜택', 'tester')
    assert (root / 'template.xlsx').read_bytes() == before_xlsx
    assert read_state(root / 'template_profile.json') == before_profiles
    assert lib.get_case(SUITE, 'BEN_0001')['sheet'] == '혜택'


def test_sheet_rename_changes_only_letter_case_without_suffix(seeded):
    import openpyxl
    lib.rename_sheet(SUITE, '혜택', 'Benefits', 'tester')
    lib.rename_sheet(SUITE, 'Benefits', 'benefits', 'tester')
    wb = openpyxl.load_workbook(lib.suite_dir(SUITE) / 'template.xlsx')
    assert 'benefits' in wb.sheetnames and 'benefits1' not in wb.sheetnames
    wb.close()
    assert lib.get_case(SUITE, 'BEN_0001')['sheet'] == 'benefits'


def test_empty_branches_are_saved_without_creating_cases(library_dir, template_xlsx):
    lib.save_template(SUITE, template_xlsx, analyze_workbook(template_xlsx))
    lib.import_cases(SUITE, ['혜택'], [], 'tester')
    lib.add_branch(SUITE, '혜택', ['직접 작성', '로그인', '오류 처리'])
    lib.add_branch(SUITE, '혜택', ['직접 작성', '로그인', '오류 처리'])
    assert lib.load_cases(SUITE) == []
    branches = lib.load_branches(SUITE)
    assert branches == [{'sheet': '혜택', 'path': ['직접 작성', '로그인', '오류 처리']}]
    tree = lib.build_tree([], branches)
    assert tree[0]['children'][0]['children'][0]['children'][0]['name'] == '오류 처리'
    assert tree[0]['count'] == 0
    for sheet, path in [('없는 시트', ['기능']), ('혜택', ['', '중분류']), ('혜택', ['기능', '', '소분류'])]:
        with pytest.raises(lib.LibraryError):
            lib.add_branch(SUITE, sheet, path)
    lib.rename_sheet(SUITE, '혜택', '사용자 작성', 'tester')
    assert lib.load_branches(SUITE)[0]['sheet'] == '사용자 작성'
    assert lib.build_tree([], lib.load_branches(SUITE))[0]['name'] == '사용자 작성'


def test_add_sheet_copies_only_format_without_sample_cases(seeded):
    import openpyxl
    lib.add_sheet(SUITE, '회원가입')
    assert len(lib.load_cases(SUITE)) == 6
    assert lib.list_suites()[0]['sheets'] == ['혜택', '홈', '회원가입']
    profile = lib.load_profiles(SUITE)['회원가입']
    assert profile.sheet == '회원가입'
    wb = openpyxl.load_workbook(lib.suite_dir(SUITE) / 'template.xlsx')
    ws = wb['회원가입']
    assert ws['G11'].value == 'Test Step'
    assert ws['G13'].value is None and ws['H13'].value is None
    assert ws['H13'].font.bold == wb['혜택']['H13'].font.bold
    assert ws.data_validations.dataValidation
    wb.close()
    for name in ['회원가입', 'bad/name', '']:
        with pytest.raises(lib.LibraryError):
            lib.add_sheet(SUITE, name)
