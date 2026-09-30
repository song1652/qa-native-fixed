"""엑셀 가져오기 성능 — LODIS 규모(시트 9개 · 약 1,800건 · 서식·병합·수식) 워크북을 서버로 끝까지 가져온다.

단계별 시간 상한을 둔다. 상한은 로컬 측정값의 약 3배라 느린 장비에서도 흔들리지 않고,
시트마다 워크북을 다시 여는 식의 회귀(미리보기 27초 → 5초로 고친 문제)는 잡는다.
실제 파일로 재려면 TC_PERF_XLSX=/경로/파일.xlsx 를 준다.
"""
from __future__ import annotations

import json
import os
import time
import urllib.parse
import urllib.request
from pathlib import Path

import openpyxl
from openpyxl.styles import Border, Font, PatternFill, Side

from tests.unit.import_studio.import_studio_test_support import dashboard_server, request_json

SHEETS = ["인증", "Wallet", "홈", "주문", "매칭", "진행중", "채팅", "마이페이지", "LODISCAN"]
ROWS_PER_SHEET = 200
BUDGET = {"preview": 15.0, "plan": 15.0, "commit": 20.0}   # 초


def lodis_like(path: Path) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin")
    for sheet in SHEETS:
        ws = wb.create_sheet(sheet)
        ws["B2"], ws["C2"], ws["C3"] = "자동계산영역", "Count", '=COUNTIF(K:M,"Pass")'
        for col, name in zip("BCDEFGHIJMOP", ["Test\nScenario ID", "Main Category", "Sub Category", "Detail Category",
                                              "TC Summuery", "Precondition", "Step", "Expected Result",
                                              "Test Result", "검증정보", "이슈번호", "이슈 및 검증 내용"]):
            ws[f"{col}10"] = name
            ws[f"{col}10"].font = Font(bold=True, color="FF0070C0")
        ws["J11"], ws["K11"], ws["L11"] = "Test Level", "Android", "iOS"
        for i in range(ROWS_PER_SHEET):
            r = 12 + i
            ws[f"B{r}"] = f'="{sheet}_"&(ROW()-11)'
            ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"] = f"대{i // 50}", f"중{i // 10}", f"소{i}"
            ws[f"F{r}"], ws[f"H{r}"], ws[f"I{r}"] = f"{sheet} 기능 {i}", "1. 실행\n2. 확인", "정상 표시"
            ws[f"J{r}"], ws[f"K{r}"], ws[f"L{r}"] = ["BAT", "Level 2", "Level 4"][i % 3], "Pass", "Fail"
            for col in "BCDEFGHIJKL":                       # 행마다 다른 서식 → 워크북 서식 표가 커진다
                cell = ws[f"{col}{r}"]
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                cell.fill = PatternFill("solid", fgColor=f"FF{(i * 37) % 256:02X}{(i * 91) % 256:02X}CC")
            if i % 20 == 0:
                ws.merge_cells(f"G{r}:G{r + 1}")
    wb.save(path)
    wb.close()
    return path


def upload(base: str, path: Path) -> tuple[int, dict]:
    request = urllib.request.Request(
        f"{base}/api/tc-library/import/preview?filename={urllib.parse.quote(path.name)}",
        data=path.read_bytes(), method="POST", headers={"Content-Type": "application/octet-stream"})
    with urllib.request.urlopen(request, timeout=120) as response:
        return response.status, json.loads(response.read())


def timed(label: str, timings: dict, fn):
    start = time.perf_counter()
    result = fn()
    timings[label] = round(time.perf_counter() - start, 2)
    return result


def import_once(base: str, path: Path, suite: str, timings: dict) -> dict:
    status, preview = timed("preview", timings, lambda: upload(base, path))
    assert status == 200, preview
    sheets = [s["name"] for s in preview["sheets"]]
    prefixes = {name: f"P{i:02d}" for i, name in enumerate(sheets)}
    payload = {"suite": suite, "sources": [{"preview_id": preview["preview_id"], "sheets": sheets,
                                            "prefixes": prefixes, "sheet_mappings": {}}]}
    status, plan = timed("plan", timings, lambda: request_json(base, "POST", "/api/tc-library/import/plan", payload))
    assert status == 200, plan
    status, result = timed("commit", timings, lambda: request_json(
        base, "POST", "/api/tc-library/import", {"run_id": plan["run_id"], "skip": [], "overwrite": []}))
    assert status == 200, result
    return {"preview": preview, "result": result}


def test_large_workbook_imports_within_budget_into_new_and_existing_suite(tmp_path):
    real = os.environ.get("TC_PERF_XLSX")
    path = Path(real) if real else lodis_like(tmp_path / "lodis_like.xlsx")
    with dashboard_server(tmp_path / "project") as base:
        first, second = {}, {}
        created = import_once(base, path, "성능확인", first)
        total = sum(s["cases"] for s in created["preview"]["sheets"])
        assert created["result"]["created"] == total > 0
        # 양식이 이미 있는 스위트에 다시 가져오기 — 다른 워크북의 서식을 옮기는 경로 (500 회귀)
        again = import_once(base, lodis_like(tmp_path / "second.xlsx") if not real else path, "성능확인", second)
        assert again["result"]["status"] == "committed"
    print(f"\n[perf] {total}건 · 새 스위트 {first} · 기존 스위트 {second}")
    for timings in (first, second):
        for step, seconds in timings.items():
            assert seconds < BUDGET[step], f"{step} {seconds}s > {BUDGET[step]}s"


def test_large_workbook_imports_into_suite_with_small_template(tmp_path):
    """기본양식처럼 서식이 적은 양식이 있는 스위트에 서식이 많은 파일을 가져와도 500이 나지 않는다 (서식 번호 회귀)."""
    small = openpyxl.Workbook()
    ws = small.active
    ws.title = "테스트케이스"
    for col, name in zip("ABCD", ["대분류", "기능", "Test Step", "Expected Result"]):
        ws[f"{col}1"] = name
    ws.append(["설치", "앱 설치", "1. 설치", "정상"])
    small.save(tmp_path / "small.xlsx")
    small.close()
    with dashboard_server(tmp_path / "project") as base:
        import_once(base, tmp_path / "small.xlsx", "작은양식", {})
        result = import_once(base, lodis_like(tmp_path / "big.xlsx"), "작은양식", {})["result"]
    assert result["created"] == len(SHEETS) * ROWS_PER_SHEET
