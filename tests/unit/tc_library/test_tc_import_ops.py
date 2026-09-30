from __future__ import annotations

import importlib
import json
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[3] / 'scripts'))
import _paths
from _tc_library import LibraryError, load_cases, patch_case


@pytest.fixture
def roots(tmp_path, monkeypatch):
    monkeypatch.setattr(_paths, 'TC_LIBRARY_DIR', tmp_path / 'library')
    return tmp_path


def upload(tmp, name='input.xlsx', feature='Login', sheet='Login', source_id='LOG_01'):
    directory = _paths.TC_LIBRARY_DIR / '_uploads'
    directory.mkdir(parents=True, exist_ok=True)
    preview_id = 'imp_' + __import__('secrets').token_hex(6)
    path = directory / (preview_id + '.xlsx')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(['ID', 'Feature', 'Steps', 'Expected', 'Tags'])
    ws.append([source_id, feature, 'Open', 'Success', 'positive'])
    wb.save(path)
    path.with_suffix('.json').write_text(json.dumps({'filename': name, 'mapping': {'header_row': 1, 'columns': {'source_tc_id': 'A', 'feature': 'B', 'steps': 'C', 'expected': 'D', 'tags': 'E'}}}))
    return {'preview_id': preview_id, 'suite': 'suite', 'sheets': [sheet], 'prefixes': {sheet: 'LOG'}}


def ops():
    # Missing module is reported as an assertion so RED demonstrates missing functionality.
    assert importlib.util.find_spec('_tc_import_ops'), 'durable library import operations are missing'
    return importlib.import_module('_tc_import_ops')


def test_import_plan_commit_repeat_and_durable_history(roots):
    operation = ops()
    body = upload(roots)
    plan = operation.plan_import(body)
    assert plan['rows'][0]['status'] == 'new'
    result = operation.commit_import({'run_id': plan['run_id']}, 'test')
    assert result['created'] == 1
    case = load_cases('suite')[0]
    assert 'auto' not in case
    repeat = operation.plan_import(upload(roots))
    assert repeat['rows'][0]['status'] == 'same'
    assert repeat['rows'][0]['case_id'] == case['case_id']
    assert operation.list_runs()[0]['run_id'] == repeat['run_id']
    assert operation.get_run(plan['run_id'])['status'] == 'committed'


def test_conflict_requires_decision_and_rollback_revisions_increase(roots):
    operation = ops()
    first = operation.plan_import(upload(roots))
    operation.commit_import({'run_id': first['run_id']}, 'test')
    case = load_cases('suite')[0]
    patch_case('suite', case['case_id'], case['rev'], {'feature': 'User edit'}, 'test')
    plan = operation.plan_import(upload(roots, feature='Imported change'))
    assert plan['rows'][0]['status'] == 'conflict'
    with pytest.raises(LibraryError, match='충돌'):
        operation.commit_import({'run_id': plan['run_id']}, 'test')
    operation.commit_import({'run_id': plan['run_id'], 'overwrite': [case['case_id']]}, 'test')
    committed = load_cases('suite')[0]
    operation.rollback_import(plan['run_id'], 'test')
    restored = load_cases('suite')[0]
    assert restored['feature'] == 'User edit'
    assert restored['rev'] > committed['rev']


def test_source_and_suite_changes_reject_commit_and_edit_rejects_rollback(roots):
    operation = ops()
    body = upload(roots)
    plan = operation.plan_import(body)
    path = _paths.TC_LIBRARY_DIR / '_uploads' / (body['preview_id'] + '.xlsx')
    path.write_bytes(path.read_bytes() + b'changed')
    with pytest.raises(LibraryError) as exc:
        operation.commit_import({'run_id': plan['run_id']}, 'test')
    assert exc.value.code == 'SOURCE_CHANGED'
    plan = operation.plan_import(upload(roots))
    operation.commit_import({'run_id': plan['run_id']}, 'test')
    case = load_cases('suite')[0]
    pending = operation.plan_import(upload(roots, feature='New'))
    patch_case('suite', case['case_id'], case['rev'], {'feature': 'Edit'}, 'test')
    with pytest.raises(LibraryError) as exc:
        operation.commit_import({'run_id': pending['run_id']}, 'test')
    assert exc.value.code == 'SUITE_CHANGED'
    with pytest.raises(LibraryError) as exc:
        operation.rollback_import(plan['run_id'], 'test')
    assert exc.value.code == 'SUITE_CHANGED'


def test_multiple_sources_preserve_existing_sheets_and_reject_same_sheet(roots):
    operation = ops()
    first = operation.plan_import(upload(roots))
    operation.commit_import({'run_id': first['run_id']}, 'test')
    a = upload(roots, name='a.xlsx', sheet='A', source_id='A_01')
    b = upload(roots, name='b.xlsx', sheet='B', source_id='B_01')
    plan = operation.plan_import({'suite': 'suite', 'sources': [a, b]})
    operation.commit_import({'run_id': plan['run_id']}, 'test')
    wb = openpyxl.load_workbook(_paths.TC_LIBRARY_DIR / 'suite/template.xlsx')
    assert wb.sheetnames == ['Login', 'A', 'B']
    wb.close()
    with pytest.raises(LibraryError) as exc:
        operation.plan_import({'suite': 'suite', 'sources': [a, upload(roots, name='other.xlsx', sheet='A')]})
    assert exc.value.code == 'DUPLICATE_SHEET'


def test_duplicate_original_id_is_error_and_skippable(roots):
    operation = ops()
    body = upload(roots)
    path = _paths.TC_LIBRARY_DIR / '_uploads' / (body['preview_id'] + '.xlsx')
    wb = openpyxl.load_workbook(path)
    wb.active.append(['LOG_01', 'Duplicate', 'Open', 'Success', 'positive'])
    wb.save(path)
    wb.close()
    plan = operation.plan_import(body)
    assert all(row['status'] == 'error' for row in plan['rows'])
    with pytest.raises(LibraryError):
        operation.commit_import({'run_id': plan['run_id']}, 'test')


def test_reimport_after_move_retains_original_identity(roots):
    operation = ops()
    first = operation.plan_import(upload(roots))
    operation.commit_import({'run_id': first['run_id']}, 'test')
    case = load_cases('suite')[0]
    patch_case('suite', case['case_id'], case['rev'], {'sheet': 'Moved'}, 'test')
    body = upload(roots, feature='Changed')
    body['prefixes'] = {'Login': 'NEW'}
    plan = operation.plan_import(body)
    assert plan['rows'][0]['case_id'] == case['case_id']
    assert plan['rows'][0]['status'] == 'conflict'


def test_interrupted_commit_rolls_forward_before_read_without_duplicate_history(roots, monkeypatch):
    operation = ops()
    first = operation.plan_import(upload(roots))
    original = operation._restore
    def interrupted(suite, snapshot):
        path = _paths.TC_LIBRARY_DIR / suite / 'cases.json'
        import base64
        operation._bytes_write(path, base64.b64decode(snapshot['cases.json']))
        raise KeyboardInterrupt('simulated process interruption')
    monkeypatch.setattr(operation, '_restore', interrupted)
    with pytest.raises(KeyboardInterrupt):
        operation.commit_import({'run_id': first['run_id']}, 'test')
    monkeypatch.setattr(operation, '_restore', original)
    assert load_cases('suite')[0]['feature'] == 'Login'
    assert operation.get_run(first['run_id'])['status'] == 'committed'
    path = _paths.TC_LIBRARY_DIR / 'suite/history.jsonl'
    entries = path.read_text().splitlines()
    load_cases('suite')
    assert path.read_text().splitlines() == entries
    assert len(entries) == 1


def test_failed_commit_restores_templates_cases_and_branches(roots, monkeypatch):
    operation = ops()
    first = operation.plan_import(upload(roots))
    operation.commit_import({'run_id': first['run_id']}, 'test')
    branches = _paths.TC_LIBRARY_DIR / 'suite/branches.json'
    branches.write_text('{"branches": ["preserve"]}')
    second = operation.plan_import(upload(roots, name='another.xlsx', sheet='B', source_id='B_01'))
    before = operation._snapshot('suite')
    original = operation._restore
    calls = []
    def failure(suite, snapshot):
        calls.append(1)
        if len(calls) == 1:
            import base64
            operation._bytes_write(_paths.TC_LIBRARY_DIR / suite / 'template.xlsx', base64.b64decode(snapshot['template.xlsx']))
            raise OSError('disk error')
        return original(suite, snapshot)
    monkeypatch.setattr(operation, '_restore', failure)
    with pytest.raises(OSError):
        operation.commit_import({'run_id': second['run_id']}, 'test')
    assert operation._snapshot('suite') == before


def test_corrupt_run_is_reported_in_listing(roots):
    operation = ops()
    first = operation.plan_import(upload(roots))
    operation._run_path(first['run_id']).write_text('{broken')
    errors = []
    assert operation.list_runs(errors) == []
    assert errors[0]['code'] == 'RUN_CORRUPT'


def test_library_import_api_persists_run_and_recovers_from_history(tmp_path):
    from tests.unit.import_studio.import_studio_test_support import dashboard_server, request_json
    from tests.unit.tc_library.test_tc_library_api import _post_bytes
    from tests.unit.tc_library.tc_fixtures import build_template_workbook
    project = tmp_path / 'project'
    (project / 'testcases').mkdir(parents=True)
    with dashboard_server(project) as base:
        status, preview = _post_bytes(base, '/api/tc-library/import/preview?filename=input.xlsx', build_template_workbook(tmp_path / 'input.xlsx').read_bytes())
        assert status == 200
        status, plan = request_json(base, 'POST', '/api/tc-library/import/plan', {'suite': 'suite', 'sources': [{'preview_id': preview['preview_id'], 'sheets': ['혜택'], 'prefixes': {'혜택': 'BEN'}}]})
        assert status == 200, plan
        run_id = plan['run_id']
        status, result = request_json(base, 'POST', '/api/tc-library/import', {'run_id': run_id})
        assert (status, result['created']) == (200, 5)
        status, listing = request_json(base, 'GET', '/api/tc-library/import/runs')
        assert status == 200
        assert listing['runs'][0]['run_id'] == run_id
        assert listing['errors'] == []
        assert 'before_snapshot' not in listing['runs'][0]
        assert 'rows' not in listing['runs'][0]
        detail = request_json(base, 'GET', '/api/tc-library/import/runs/' + run_id)[1]
        assert 'before_snapshot' not in detail and 'journal' not in detail
        assert 'rows' in detail
        assert request_json(base, 'GET', '/api/tc-library/import/runs/' + run_id)[1]['status'] == 'committed'
        status, result = request_json(base, 'POST', '/api/tc-library/import/runs/' + run_id + '/rollback', {})
        assert (status, result['status']) == (200, 'rolled_back')
        assert load_cases('suite') == []
        tombstones = load_cases('suite', include_deleted=True)
        assert all(case['deleted'] and case['rev'] == 2 for case in tombstones)


def test_case_writer_waits_for_whole_suite_transaction(roots):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event
    from _tc_library import suite_lock
    operation = ops()
    first = operation.plan_import(upload(roots))
    operation.commit_import({'run_id': first['run_id']}, 'test')
    case = load_cases('suite')[0]
    started, finished = Event(), Event()
    def writer():
        started.set()
        patch_case('suite', case['case_id'], case['rev'], {'feature': 'Concurrent edit'}, 'thread')
        finished.set()
    with ThreadPoolExecutor(max_workers=1) as pool:
        with suite_lock('suite'):
            with suite_lock('suite'):
                future = pool.submit(writer)
                assert started.wait(1)
                assert not finished.wait(0.05)
        future.result(timeout=2)
    assert load_cases('suite')[0]['feature'] == 'Concurrent edit'


def test_import_preserves_auxiliary_sheets_and_exposes_only_public_history(roots):
    operation = ops()
    body = upload(roots)
    path = _paths.TC_LIBRARY_DIR / '_uploads' / (body['preview_id'] + '.xlsx')
    wb = openpyxl.load_workbook(path)
    history = wb.create_sheet('History')
    history['A1'] = 'Keep this original history'
    wb.save(path)
    wb.close()
    plan = operation.plan_import(body)
    operation.commit_import({'run_id': plan['run_id']}, 'test')
    wb = openpyxl.load_workbook(_paths.TC_LIBRARY_DIR / 'suite/template.xlsx')
    assert wb['History']['A1'].value == 'Keep this original history'
    wb.close()
    result = operation.public_run(operation.get_run(plan['run_id']))
    assert not {'before_snapshot', 'journal', 'before_hash', 'after_hash'}.intersection(result)
    assert 'profiles' not in result['sources'][0]
    assert 'rows' in result
    assert 'rows' not in operation.public_run(operation.get_run(plan['run_id']), summary_only=True)


def test_existing_original_id_from_another_file_is_not_silently_duplicated(roots):
    operation = ops()
    first = operation.plan_import(upload(roots))
    operation.commit_import({'run_id': first['run_id']}, 'test')
    plan = operation.plan_import(upload(roots, name='different.xlsx'))
    assert plan['rows'][0]['status'] == 'error'
    assert len(load_cases('suite')) == 1


@pytest.mark.parametrize('bad_source', [None, {'preview_id': 'not-valid'}, {'preview_id': 'imp_000000000000', 'sheets': 'Login'}])
def test_invalid_source_payload_reports_library_error(roots, bad_source):
    operation = ops()
    with pytest.raises(LibraryError):
        operation.plan_import({'suite': 'suite', 'sources': [bad_source]})


@pytest.mark.parametrize('sheet_default, auto_value', [
    ('Y-app', None),
    ('', None),
    (None, None),
    ('Y-web', 'N'),
    ('Y-web', ''),
])
def test_retired_auto_mapping_options_are_ignored(roots, sheet_default, auto_value):
    operation = ops()
    body = upload(roots)
    path = _paths.TC_LIBRARY_DIR / '_uploads' / (body['preview_id'] + '.xlsx')
    mapping = json.loads(path.with_suffix('.json').read_text())['mapping']
    if sheet_default is not None:
        mapping['default_auto'] = sheet_default
    if auto_value is not None:
        mapping['columns']['auto'] = 'F'
        wb = openpyxl.load_workbook(path)
        wb.active['F1'] = 'AUTO'
        wb.active['F2'] = auto_value
        wb.save(path)
        wb.close()
    body['sheet_mappings'] = {'Login': mapping}
    plan = operation.plan_import(body)
    assert 'auto' not in plan['rows'][0]['after']
    operation.commit_import({'run_id': plan['run_id']}, 'test')
    assert 'auto' not in load_cases('suite')[0]


def test_retired_auto_default_does_not_restrict_import(roots):
    operation = ops()
    body = upload(roots)
    path = _paths.TC_LIBRARY_DIR / '_uploads' / (body['preview_id'] + '.xlsx')
    mapping = json.loads(path.with_suffix('.json').read_text())['mapping']
    mapping['default_auto'] = 'unsupported'
    body['sheet_mappings'] = {'Login': mapping}
    plan = operation.plan_import(body)
    assert plan['rows'][0]['status'] == 'new'
    assert 'auto' not in plan['rows'][0]['after']


def test_imported_cases_are_review_locked_and_reimport_keeps_human_review(roots):
    operation = ops()
    operation.commit_import({'run_id': operation.plan_import(upload(roots))['run_id']}, 'test')
    case = load_cases('suite')[0]
    assert case['review_source'] == 'import' and case['status'] == 'approved'
    with pytest.raises(LibraryError) as exc:                              # 가져온 케이스는 검토 상태 잠금
        patch_case('suite', case['case_id'], case['rev'], {'status': 'rejected'}, 'test')
    assert exc.value.code == 'IMPORTED_REVIEW_LOCKED'
    patch_case('suite', case['case_id'], case['rev'], {'feature': '편집은 됨'}, 'test')   # 내용 편집은 가능
    # 사람이 검토한 케이스로 바뀐 뒤 같은 원본 ID로 다시 가져와도 '가져옴'으로 되돌아가지 않는다
    from _tc_library import set_review_source
    set_review_source('suite', case['case_id'], 'human')
    plan = operation.plan_import(upload(roots, feature='엑셀 수정'))
    operation.commit_import({'run_id': plan['run_id'], 'overwrite': [case['case_id']]}, 'test')
    assert load_cases('suite')[0]['review_source'] == 'human'
