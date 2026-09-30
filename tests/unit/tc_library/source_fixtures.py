"""소스 테스트용 문서 생성기 — PDF는 손으로 만든 최소 파일, DOCX는 python-docx로 만든다."""
from __future__ import annotations

import io


def make_pdf(pages: list[str]) -> bytes:
    """ASCII 텍스트 페이지로 된 최소 PDF. 빈 문자열이면 텍스트 없는 쪽(스캔본 흉내)."""
    objects: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>", b""]
    font_id = 3
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    kids = []
    for text in pages:
        stream = (f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET" if text else "").encode("latin-1")
        objects.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream))
        content_id = len(objects)
        objects.append(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                       b"/Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>" % (font_id, content_id))
        kids.append(len(objects))
    objects[1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (
        b" ".join(b"%d 0 R" % k for k in kids), len(kids))
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n%s\nendobj\n" % (number, body))
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def make_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("배너 롤링 규칙", level=1)
    doc.add_paragraph("배너는 3초마다 자동으로 다음 배너로 이동한다.")
    doc.add_heading("배너 선택 동작", level=2)
    doc.add_paragraph("배너를 누르면 설정된 링크로 이동한다.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "조건", "결과"
    table.cell(1, 0).text, table.cell(1, 1).text = "배너 1개", "인디케이터 미노출"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


PRD_MD = """# 8.6.0 혜택 탭 상단 배너 개편

## 배너 롤링 규칙
배너는 최대 5개까지 등록하며 등록 순서대로 노출한다.
배너는 3초마다 자동으로 다음 배너로 이동한다.

## 배너 선택 동작
배너 선택 시 배너에 설정된 링크로 이동한다.
이미지를 불러오지 못하면 기본 이미지와 "혜택을 준비하고 있어요" 문구를 노출한다.
"""
