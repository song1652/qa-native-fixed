"""야핏무브 Full TC 구조를 흉내 낸 작은 워크북 (실제 파일은 저장소에 넣지 않는다)."""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = ["No.", "대분류", "중분류", "소분류", "기능", "사전 조건", "Test Step",
           "Expected Result", "우선순위", "AUTO", "환경", None, "기타"]
NO_FORMULA = '=IF(H{r}<>"",ROW(B{r})-12, "")'

BENEFIT_ROWS = [
    # B, C, D, E, F, G, H, I, J, K, L, M
    ("혜택 탭", None, None, "혜택 탭 버튼", None, "1. 앱 실행\n2. 혜택 탭 선택",
     "혜택 탭 화면으로 진입된다.", "P0", None, None, None, None),
    (None, "상단 배너", None, "배너 스크롤", None, "1. 혜택 탭 선택",
     "상단에 광고 배너가 가로 스크롤 동작되어 노출된다.\n- 수동, 자동 스크롤",
     None, None, "Pass", "Pass", None),
    (None, None, None, "광고 배너", None, "1. 혜택 탭 선택\n2. 광고 배너 선택",
     "해당 상세 페이지로 진입된다.", None, None, "Fail", "Pass", None),
    (None, "신규회원 한정 혜택", "돈불리기", "진입 불가", "- D+0 가입일자",
     "1. 혜택 탭 선택\n2. 돈불리기 선택",
     "진입 불가 안내 팝업이 노출된다.\n- 내일부터 참여할 수 있어요",
     None, None, None, None, "돈불리기 정책 변경"),
    (None, None, None, None, "- D+13 가입일자", "1. 혜택 탭 선택\n2. 돈불리기 선택",
     "진입 불가 안내 팝업이 노출된다.\n- 마지막 날이에요", "P2", None, "NA", None, None),
]
HOME_ROWS = [
    ("걷고 받기 탭", None, None, "코호트 별 홈 화면", "- 가입일(D day) ~ D+3", "1. 앱 실행",
     "가입일(D day) ~ D+3 상태 홈 화면이 노출된다.", "P1", None, None, None, None),
]


def _summary_block(ws) -> None:
    ws["I2"], ws["J2"] = "구분", "COUNT"
    ws["J3"], ws["K3"], ws["L3"], ws["M3"] = "And", "iOS", "And", "iOS"
    ws["I4"] = "Pass"
    ws["J4"], ws["K4"] = "=COUNTIF(#REF!,I4)", "=COUNTIF(#REF!,I4)"
    ws["I9"] = "Total Case"
    ws["J9"] = "=IF($L3=0,0,COUNTA($A$13:$A$390))"
    ws["K9"] = "=IF($L3=0,0,COUNTA($A$13:$A$390))"


def _tc_sheet(wb, title: str, rows: list[tuple]) -> None:
    ws = wb.create_sheet(title)
    _summary_block(ws)
    for col, header in enumerate(HEADERS, 1):
        ws.cell(11, col).value = header
    ws["K12"], ws["L12"] = "And", "iOS"
    for i, row in enumerate(rows):
        r = 13 + i
        ws.cell(r, 1).value = NO_FORMULA.format(r=r)
        for col, value in enumerate(row, 2):
            ws.cell(r, col).value = value
    ws["H13"].font = Font(bold=True)
    for formula, ref in (('"P0,P1,P2"', "I13:I40"), ('"AUTO"', "J13:J40"),
                         ('"Pass,Fail,NT,NA"', "K13:L40")):
        dv = DataValidation(type="list", formula1=formula)
        dv.add(ref)
        ws.add_data_validation(dv)


def build_template_workbook(path: Path) -> Path:
    wb = openpyxl.Workbook()
    history = wb.active
    history.title = "History"
    history["B1"] = "Full TC 관리"
    history["B2"], history["C2"] = "날짜", "반영 내용"
    history["B3"], history["C3"] = "25.09.11", "초기 작성"
    _tc_sheet(wb, "혜택", BENEFIT_ROWS)
    wb["혜택"].merge_cells("D16:D17")
    _tc_sheet(wb, "홈", HOME_ROWS)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path
