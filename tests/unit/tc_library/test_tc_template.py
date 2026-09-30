from __future__ import annotations

from _tc_template import TemplateProfile, analyze_workbook


def test_header_row_skips_summary_block_and_finds_sub_header(template_xlsx):
    profiles = analyze_workbook(template_xlsx)

    assert list(profiles) == ["혜택", "홈"]          # History는 TC 시트가 아니다
    benefit = profiles["혜택"]
    assert benefit.header_row == 11
    assert benefit.data_start_row == 13
    assert benefit.result_columns == {"And": 11, "iOS": 12}
    assert benefit.columns["l1"] == 2 and benefit.columns["note"] == 13


def test_dropdowns_and_no_formula_are_captured(template_xlsx):
    benefit = analyze_workbook(template_xlsx)["혜택"]

    assert benefit.validations == {
        "priority": ["P0", "P1", "P2"],
        "execution_result": ["Pass", "Fail", "NT", "NA"],
    }
    assert benefit.no_formula == '=IF(H{r}<>"",ROW(B{r})-12, "")'
    assert "우선순위 드롭다운을 내보낼 때 P0~P3으로 넓힙니다" in benefit.warnings


def test_profile_round_trips_through_dict(template_xlsx):
    benefit = analyze_workbook(template_xlsx)["혜택"]
    assert TemplateProfile.from_dict(benefit.to_dict()) == benefit


def test_english_category_headers_are_recognized(tmp_path):
    """LODIS 양식: 요약 표 아래 10행이 Main Category / TC Summuery(오타) / Step 헤더 — 자동 인식돼야 한다."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "인증"
    ws["B2"], ws["C2"] = "자동계산영역", "Count"
    for col, name in zip("BCDEFGHI", ["Test\nScenario ID", "Main Category", "Sub Category", "Detail Category",
                                      "TC Summuery", "Precondition", "Step", "Expected Result"]):
        ws[f"{col}10"] = name
    ws["C12"], ws["F12"], ws["H12"], ws["I12"] = "설치", "앱 설치 확인", "1. 설치", "정상 설치"
    path = tmp_path / "lodis.xlsx"
    wb.save(path)
    profile = analyze_workbook(path)["인증"]
    assert profile.header_row == 10
    assert {"l1", "l2", "l3", "feature", "precondition", "steps", "expected"} <= set(profile.columns)


def test_test_level_sub_header_becomes_priority_and_formula_id_is_not_stored(tmp_path):
    """결과 아래 칸의 Test Level은 우선순위 열이고, 수식 원본 ID는 수식 문자열로 저장되지 않는다."""
    import openpyxl
    from _tc_xlsx_import import import_workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "인증"
    for col, name in zip("BCFHIJ", ["Test Scenario ID", "Main Category", "TC Summary", "Step", "Expected Result", "Test Result"]):
        ws[f"{col}10"] = name
    ws["J11"], ws["K11"] = "Test Level", "Android"
    rows = [("BAT", "P0"), ("Level 2", "P1"), ("Level 4", "P3")]
    for i, (level, _) in enumerate(rows, 12):
        ws[f"B{i}"] = f'="인증_"&(ROW()-11)'
        ws[f"C{i}"], ws[f"F{i}"], ws[f"H{i}"], ws[f"I{i}"] = "설치", f"기능{i}", "1. 실행", "정상"
        ws[f"J{i}"], ws[f"K{i}"] = level, "Pass"
    path = tmp_path / "lodis.xlsx"
    wb.save(path)
    profiles = analyze_workbook(path)
    profile = profiles["인증"]
    assert profile.data_start_row == 12
    assert "Test Level" not in profile.result_columns and profile.columns["priority"] == 10
    cases = import_workbook(path, profiles, ["인증"], {})
    assert [c["priority"] for c in cases] == [p for _, p in rows]
    assert all(not c["source_tc_id"].startswith("=") for c in cases)
    assert {c["execution_result"] for c in cases} == {"pass"}
