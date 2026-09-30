from __future__ import annotations

import hashlib
import json

import openpyxl

from _tc_model import EDITABLE_FIELDS, new_case, validate_case
from _tc_template import analyze_workbook, mapping_from_import_profile
from _tc_xlsx_import import import_workbook
from _tc_xlsx_export import export_workbook, verify_export


def test_new_case_and_validation_ignore_retired_auto():
    case = new_case(case_id='A_0001', sheet='Cases', path=['Login'], feature='Login', steps=['Open'], expected='Visible', priority='P0', auto='legacy unknown')
    assert 'auto' not in case and 'auto' not in EDITABLE_FIELDS
    legacy = {**case, 'auto': {'obsolete': True}}
    assert not any(issue['code'] == 'AUTO_INVALID' for issue in validate_case(legacy))


def test_import_no_longer_maps_or_populates_auto(template_xlsx):
    profiles = analyze_workbook(template_xlsx)
    assert all('auto' not in profile.columns for profile in profiles.values())
    assert mapping_from_import_profile({'auto': 'F열', 'title': 'B열'}) == {'feature': 'B'}
    cases = import_workbook(template_xlsx, profiles, ['혜택'], {'혜택': 'BEN'})
    assert len(cases) == 5 and all('auto' not in case for case in cases)


def test_xlsx_removes_auto_column_preserves_neighbors_formulas_and_original(template_xlsx, tmp_path):
    original = hashlib.sha256(template_xlsx.read_bytes()).hexdigest()
    profiles = analyze_workbook(template_xlsx)
    cases = import_workbook(template_xlsx, profiles, ['혜택', '홈'], {'혜택': 'BEN', '홈': 'HOME'})
    by_sheet = {name: [case for case in cases if case['sheet'] == name] for name in ('혜택', '홈')}
    out = tmp_path / 'no_auto.xlsx'
    export_workbook(template_xlsx, profiles, by_sheet, out)
    wb = openpyxl.load_workbook(out)
    ws = wb['혜택']
    assert 'AUTO' not in [cell.value for cell in ws[11]]
    assert ws['I13'].value == 'P0'
    assert ws['J15'].value == ws['K15'].value == 'Fail'
    assert ws['L16'].value == '돈불리기 정책 변경\nid:BEN_0004'
    assert ws['J4'].value == '=COUNTIF($J$13:$J$17,I4)'
    dvs = {validation.formula1: str(validation.sqref) for validation in ws.data_validations.dataValidation}
    assert dvs['"Pass,Fail,NT,NA"'] == 'J13:K17'
    wb.close()
    assert hashlib.sha256(template_xlsx.read_bytes()).hexdigest() == original
    assert all(check['level'] == 'ok' for check in verify_export(out, profiles, by_sheet))


def test_library_and_profile_api_hide_legacy_auto_fields(tmp_path):
    from tests.unit.import_studio.import_studio_test_support import dashboard_server, request_json
    with dashboard_server(tmp_path / 'project') as url:
        import _paths
        root = _paths.TC_LIBRARY_DIR / 'suite'
        root.mkdir(parents=True)
        case = new_case(case_id='A_0001', sheet='Cases', path=['Login'], feature='Login', steps=['Open'], expected='Visible', priority='P0')
        case['auto'] = 'obsolete'
        (root / 'cases.json').write_text(json.dumps({'suite': 'suite', 'sheets': ['Cases'], 'cases': [case]}))
        profile = {'id': 'old', 'name': 'old', 'default_auto': 'Y-web', 'mappings': {'title': 'B열', 'steps': 'C열', 'expected': 'D열', 'auto': 'F열'}}
        _paths.IMPORT_PROFILES_PATH.write_text(json.dumps({'profiles': [profile]}))
        listing = request_json(url, 'GET', '/api/tc-library/suite?auto=Y-web')[1]
        assert listing['total'] == 1 and 'auto' not in listing['items'][0]
        detail = request_json(url, 'GET', '/api/tc-library/suite/cases/A_0001')[1]['case']
        assert 'auto' not in detail
        public_profile = request_json(url, 'GET', '/api/tc-library/import/mapping-profiles')[1]['profiles'][0]
        assert 'auto' not in public_profile['columns'] and 'default_auto' not in public_profile['mapping']
        assert json.loads((root / 'cases.json').read_text())['cases'][0]['auto'] == 'obsolete'
        assert request_json(url, 'GET', '/api/import/profiles')[1]['profiles'][0]['mappings']['auto'] == 'F열'


def test_xlsx_removes_auto_summary_and_dependent_ratios_keeps_other_summary(template_xlsx, tmp_path):
    wb = openpyxl.load_workbook(template_xlsx)
    ws = wb['혜택']
    ws['N2'] = 'AUTO'
    ws['N3'] = 'Y(web)'
    ws['N4'] = 'Y(app)'
    ws['N5'] = 'N'
    ws['N7'] = '=COUNTIF($J$13:$J$390,"Y(web)")'
    ws['N8'] = '=COUNTIF($J$13:$J$390,"Y(app)")'
    ws['N9'] = '=SUM(N7:N8)/J9'
    ws['O9'] = '=N9*100'
    wb.save(template_xlsx)
    wb.close()
    profiles = analyze_workbook(template_xlsx)
    cases = import_workbook(template_xlsx, profiles, ['혜택'], {'혜택': 'BEN'})
    out = tmp_path / 'no_auto_summary.xlsx'
    export_workbook(template_xlsx, profiles, {'혜택': cases}, out)
    wb = openpyxl.load_workbook(out)
    ws = wb['혜택']
    assert all(ws[f'N{row}'].value is None for row in (2, 3, 4, 5, 7, 8, 9))
    assert ws['O9'].value is None
    assert ws['I4'].value == 'Pass'
    assert ws['J4'].value == '=COUNTIF($J$13:$J$17,I4)'
    assert not any('AUTO' in str(cell.value) or '#REF!' in str(cell.value) for row in ws for cell in row)
    wb.close()
    assert all(check['level'] == 'ok' for check in verify_export(out, profiles, {'혜택': cases}))


def test_xlsx_relocates_absolute_relative_and_cross_sheet_whole_column_formulas(template_xlsx, tmp_path):
    wb = openpyxl.load_workbook(template_xlsx)
    ws = wb['혜택']
    ws['P4'] = '=COUNTIF($K:$K,"Pass")'
    ws['P5'] = '=SUM(K:K)'
    ws['P6'] = '=SUM(K:$L)'
    wb['History']['D1'] = "=COUNTIF('혜택'!$K:$K,\"Pass\")"
    wb.save(template_xlsx)
    wb.close()
    profiles = analyze_workbook(template_xlsx)
    cases = import_workbook(template_xlsx, profiles, ['혜택'], {'혜택': 'BEN'})
    out = tmp_path / 'whole_columns.xlsx'
    export_workbook(template_xlsx, profiles, {'혜택': cases}, out)
    wb = openpyxl.load_workbook(out)
    ws = wb['혜택']
    assert ws['P4'].value == '=COUNTIF($J:$J,"Pass")'
    assert ws['P5'].value == '=SUM(J:J)'
    assert ws['P6'].value == '=SUM(J:$K)'
    assert wb['History']['D1'].value == "=COUNTIF('혜택'!$J:$J,\"Pass\")"
    wb.close()
