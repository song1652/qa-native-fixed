from __future__ import annotations

from datetime import date

import openpyxl

from _tc_model import new_case
from _tc_template import analyze_workbook
from _tc_xlsx_export import export_workbook, verify_export
from _tc_xlsx_import import import_workbook


def _load(template_xlsx):
    profiles = analyze_workbook(template_xlsx)
    cases = import_workbook(template_xlsx, profiles, ["혜택", "홈"], {"혜택": "BEN", "홈": "HOME"})
    by_sheet: dict[str, list[dict]] = {}
    for case in cases:
        by_sheet.setdefault(case["sheet"], []).append(case)
    return profiles, by_sheet


def test_export_rewrites_rows_formulas_merges_and_dropdowns(template_xlsx, tmp_path):
    profiles, by_sheet = _load(template_xlsx)
    by_sheet["혜택"][4]["priority"] = "P3"
    out = tmp_path / "out.xlsx"
    export_workbook(template_xlsx, profiles, by_sheet, out, history_note="8.6.0 반영",
                    today=date(2026, 9, 30))

    wb = openpyxl.load_workbook(out)
    ws = wb["혜택"]
    assert ws["A17"].value == '=IF(H17<>"",ROW(B17)-12, "")'
    assert ws["J4"].value == "=COUNTIF($J$13:$J$17,I4)"               # #REF! 복구
    assert ws["J9"].value == "=IF($L3=0,0,COUNTA($A$13:$A$17))"       # 고정 범위 → 실제 범위
    assert {str(r) for r in ws.merged_cells.ranges} == {"B13:B17", "C14:C15", "C16:C17", "D16:D17", "E16:E17"}
    dvs = {dv.formula1: str(dv.sqref) for dv in ws.data_validations.dataValidation}
    assert dvs['"P0,P1,P2,P3"'] == "I13:I17"
    assert dvs['"Pass,Fail,NT,NA"'] == "J13:K17"
    assert ws["I17"].value == "P3"
    assert ws["J15"].value == ws["K15"].value == "Fail"
    assert ws["H17"].font.b is True                                  # 13행 스타일 복사
    assert ws["L16"].value == "돈불리기 정책 변경\nid:BEN_0004"
    assert [c.value for c in wb["History"][4]][1:3] == ["26.09.30", "8.6.0 반영"]

def test_export_then_reimport_is_lossless(template_xlsx, tmp_path):
    profiles, by_sheet = _load(template_xlsx)
    by_sheet["혜택"].append(new_case(case_id="BEN_0006", sheet="혜택",
                                     path=["혜택 탭", "상단 배너", ""], feature="새 기능"))
    out = tmp_path / "out.xlsx"
    export_workbook(template_xlsx, profiles, by_sheet, out)

    checks = verify_export(out, profiles, by_sheet)
    assert {c["level"] for c in checks} == {"ok"}, checks
    again = import_workbook(out, analyze_workbook(out), ["혜택", "홈"], {"혜택": "ZZZ", "홈": "ZZZ"})
    assert [c["case_id"] for c in again] == [c["case_id"] for s in by_sheet.values() for c in s]


def test_verify_export_reports_mismatch(template_xlsx, tmp_path):
    profiles, by_sheet = _load(template_xlsx)
    out = tmp_path / "out.xlsx"
    export_workbook(template_xlsx, profiles, by_sheet, out)
    by_sheet["홈"][0]["feature"] = "다른 이름"

    checks = verify_export(out, profiles, by_sheet)
    assert ("error", "ROUNDTRIP") in {(c["level"], c["code"]) for c in checks}
